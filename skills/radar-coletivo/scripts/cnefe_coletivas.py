#!/usr/bin/env python3
"""
CNEFE IBGE 2022 — Análise de Ligações Coletivas
=================================================
Normaliza complementos, classifica coletivas, gera Excel formatado.

Uso:
  python cnefe_coletivas.py --input 4211900_PALHOCA.csv
  python cnefe_coletivas.py --input CNEFE_2022_SC.csv.gz --municipio 4211900
  python cnefe_coletivas.py --input 4211900_PALHOCA.csv --output minha_analise.xlsx
"""

import argparse
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ╔═══════════════════════════════════════════════════════════════╗
# ║  CONSTANTES E DICIONÁRIOS                                    ║
# ╚═══════════════════════════════════════════════════════════════╝

# v4.3: LAYER_MAP/CAMADAS têm fonte ÚNICA em radar_utils (a cópia local havia
# divergido — perdia ESQ/DIR/MEIO/LATERAL da camada POSICAO).
from radar_utils import LAYER_MAP, CAMADAS, strip_accents

QUALIDADE_GEO = {1: 'VALIDADA', 2: 'VALIDADA', 3: 'ESTIMADA', 4: 'BAIXA', 5: 'BAIXA', 6: 'BAIXA'}
TIPO_ESP = {101: 'Casa', 102: 'Casa vila/cond', 103: 'Apartamento', 104: 'Outros'}
ESPECIE = {
    1: 'Residencial', 2: 'Dom. coletivo', 3: 'Agropecuário',
    4: 'Ensino', 5: 'Saúde', 6: 'Outras finalid.',
    7: 'Construção', 8: 'Religioso',
}

TIPO_DESC = {
    'COND_VERT_MULTIBLOCOS': 'Condomínio vertical com múltiplos blocos/torres',
    'COND_HORIZ_MULTIBLOCOS': 'Condomínio horizontal com múltiplos blocos/vias',
    'COND_MULTIBLOCOS': 'Condomínio com múltiplos blocos (tipo misto)',
    'COND_VIA_INT': 'Condomínio com via interna (alameda/rua interna)',
    'EMPREEND_NOMEADO': 'Empreendimento nomeado (prédio/edifício/galeria)',
    'VERT_APTO': 'Edifício de apartamentos',
    'VERT_COMPL': 'Edifício com unidades (sala/loja/quitinete)',
    'VERT_BLOCO': 'Edifício com blocos identificados',
    'HORIZ_VILA_COND': 'Vila ou condomínio horizontal (casas)',
    'HORIZ_MORADIAS': 'Múltiplas moradias no mesmo endereço (casa/lote/sobrado)',
    'HORIZ_SUBDIV': 'Subdivisão horizontal (fundos/frente/lado/anexo)',
    'MULT_ESTAB': 'Múltiplos estabelecimentos no endereço (indicador oficial)',
    'AGRUP_END': 'Agrupamento por endereço (sem complemento específico)',
    'APTO_ISOLADO': 'Apartamento isolado (único no endereço)',
    'AGRUP_TEXTO': 'Agrupamento textual (coordenada estimada/baixa)',
    'DUPLA_END': 'Apenas 2 registros no endereço',
    'NAO_COLETIVA': 'Não coletiva (registro individual)',
}
CONF_DESC = {
    'ALTA': 'Alta — tipologia confirmada + coordenada validada',
    'MEDIA': 'Média — evidência parcial ou coordenada estimada',
    'BAIXA': 'Baixa — agrupamento sem âncora estrutural confiável',
    'NAO_COLETIVA': 'Não coletiva',
}

# ── classificação lexical de atividade: regras de keyword match
_LEXICO_ATIVIDADE = [
    # (setor, atividade_prefix, keywords_OR, match_type)
    # match_type: 'exact' = d == kw, 'contains' = kw in d
    ('ALIMENTACAO', 'Bar', ['BAR'], 'exact_or_prefix'),
    ('ALIMENTACAO', 'Restaurante', ['RESTAURANTE']),
    ('ALIMENTACAO', 'Lanchonete', ['LANCHONETE']),
    ('ALIMENTACAO', 'Padaria', ['PADARIA', 'PANIFICADORA']),
    ('ALIMENTACAO', 'Pizzaria', ['PIZZARIA']),
    ('ALIMENTACAO', 'Hamburgueria', ['HAMBURGUERIA']),
    ('ALIMENTACAO', 'Sorveteria', ['SORVETERIA']),
    ('ALIMENTACAO', 'Confeitaria', ['CONFEITARIA']),
    ('ALIMENTACAO', 'Cafeteria', ['CAFE', 'CAFETERIA']),
    ('ALIMENTACAO', 'Choperia/bar', ['CHOPERIA']),
    ('ALIMENTACAO', 'Açaí/sucos', ['ACAI']),
    ('ALIMENTACAO', 'Mercado/mercearia', ['MERCADO', 'MERCEARIA', 'MINIMERCADO']),
    ('ALIMENTACAO', 'Açougue/frigorífico', ['ACOUGUE', 'FRIGORIFICO', 'ABATEDOURO']),
    ('ALIMENTACAO', 'Peixaria', ['PEIXARIA']),
    ('INDUSTRIAL', 'Fábrica', ['FABRICA']),
    ('INDUSTRIAL', 'Indústria', ['INDUSTRIA']),
    ('INDUSTRIAL', 'Metalúrgica', ['METALURGICA']),
    ('INDUSTRIAL', 'Usina', ['USINA']),
    ('INDUSTRIAL', 'Galpão/depósito industrial', ['GALPAO']),
    ('INDUSTRIAL', 'Depósito/armazém', ['DEPOSITO']),
    ('INDUSTRIAL', 'Armazém', ['ARMAZEM']),
    ('MANUFATURA', 'Marcenaria', ['MARCENARIA']),
    ('MANUFATURA', 'Serralheria', ['SERRALHERIA']),
    ('MANUFATURA', 'Vidraçaria', ['VIDRACARIA']),
    ('MANUFATURA', 'Funilaria', ['FUNILARIA']),
    ('MANUFATURA', 'Estofaria/tapeçaria', ['ESTOFARIA', 'TAPECAR']),
    ('MANUFATURA', 'Costura/alfaiataria', ['COSTURA', 'ALFAIAT']),
    ('MANUFATURA', 'Borracharia', ['BORRACHARIA']),
    ('MANUFATURA', 'Marmoraria', ['MARMORARIA']),
    ('MANUFATURA', 'Gráfica', ['GRAFICA']),
    ('SERVICO_AUTO', 'Oficina mecânica', ['OFICINA MECANICA', 'OFICINA AUTOMOTIVA']),
    ('SERVICO_AUTO', 'Oficina', ['OFICINA']),
    # v4.3: LAVANDERIA antes do lava-rápido (o substring 'LAVA' capturava lavanderia)
    ('COMERCIO_SERVICO', 'Lavanderia', ['LAVANDERIA']),
    ('SERVICO_AUTO', 'Lava-rápido', ['LAVA JATO', 'LAVA-JATO', 'LAVA RAPIDO', 'LAVA-RAPIDO',
                                     'LAVACAO', 'LAVA CAR', 'LAVAGEM']),
    ('SERVICO_AUTO', 'Auto elétrica', ['AUTO ELETRICA']),
    ('SERVICO_AUTO', 'Estacionamento', ['ESTACIONAMENTO']),
    ('COMERCIO_VAREJO', 'Vestuário', ['LOJA DE ROUPA', 'BOUTIQUE', 'CONFECCAO']),
    ('COMERCIO_VAREJO', 'Material de construção', ['MATERIAL DE CONSTRUCAO', 'MADEIREIRA']),
    ('COMERCIO_VAREJO', 'Farmácia/drogaria', ['FARMACIA', 'DROGARIA']),
    # v4.3: POSTO antes de PET ('PET' substring capturava 'POSTO PETROBRAS');
    # POSTO DE SAUDE protegido; AUTO POSTO coberto; PET restrito a
    # exact/prefix + PETSHOP ('PET' solto capturava 'TAPETE')
    ('SAUDE', 'Posto de saúde', ['POSTO DE SAUDE', 'POSTO SAUDE']),
    ('COMERCIO_VAREJO', 'Posto de combustível', ['AUTO POSTO', 'AUTOPOSTO',
                                                 'POSTO DE COMBUSTIVEL', 'POSTO DE GASOLINA']),
    ('COMERCIO_VAREJO', 'Posto de combustível', ['POSTO'], 'exact_or_prefix'),
    ('COMERCIO_VAREJO', 'Pet shop', ['PET SHOP', 'PETSHOP']),
    ('COMERCIO_VAREJO', 'Pet shop', ['PET'], 'exact_or_prefix'),
    ('COMERCIO_VAREJO', 'Brechó', ['BRECHO']),
    ('COMERCIO_VAREJO', 'Papelaria', ['PAPELARIA']),
    ('COMERCIO_VAREJO', 'Conveniência', ['CONVENIENCIA']),
    ('BELEZA_ESTETICA', 'Salão de beleza', ['SALAO DE BELEZA', 'SALAO']),
    ('BELEZA_ESTETICA', 'Barbearia', ['BARBEARIA']),
    ('BELEZA_ESTETICA', 'Estética', ['ESTETICA']),
    ('SERVICO_PROF', 'Escritório', ['ESCRITORIO']),
    ('SERVICO_PROF', 'Consultório', ['CONSULTORIO']),
    ('SERVICO_PROF', 'Imobiliária', ['IMOBILIARIA']),
    ('SERVICO_PROF', 'Contabilidade', ['CONTABIL']),
    ('SERVICO_PROF', 'Advocacia', ['ADVOCACIA', 'ADVOGAD']),
    ('SERVICO_PROF', 'Cartório', ['CARTORIO']),
    ('SAUDE', 'Clínica', ['CLINICA']),
    ('SAUDE', 'Laboratório', ['LABORATORIO']),
    ('SAUDE', 'Academia', ['ACADEMIA']),
    ('SAUDE', 'Veterinária', ['VETERINAR']),
    ('EDUCACAO', 'Escola', ['ESCOLA', 'COLEGIO']),
    ('EDUCACAO', 'Creche', ['CRECHE']),
    ('EDUCACAO', 'Autoescola', ['AUTO ESCOLA', 'AUTOESCOLA']),
    ('EDUCACAO', 'Curso', ['CURSO']),
    ('HOSPEDAGEM', 'Pousada', ['POUSADA']),
    ('HOSPEDAGEM', 'Hotel', ['HOTEL']),
    ('HOSPEDAGEM', 'Hostel', ['HOSTEL']),
    ('HOSPEDAGEM', 'Motel', ['MOTEL']),
    ('ASSOCIACAO_ONG', 'Associação/cooperativa', ['ASSOCIACAO', 'ONG', 'SINDICATO', 'COOPERATIVA']),
    ('LAZER', 'Rancho de pesca', ['RANCHO']),
    ('LAZER', 'Espaço esportivo', ['CAMPO DE FUTEBOL', 'QUADRA ESPORTIVA']),
    ('VAGO', 'Imóvel vago/desocupado', ['VAGO', 'VAGA', 'DESOCUPADO', 'ABANDONADO', 'FECHADO']),
    ('VAGO', 'Sem identificação', ['SEM NOME', 'NI', 'SEM DENOMINACAO', 'SEM IDENTIFICACAO']),
    ('DIVERSOS', 'Garagem', ['GARAGEM']),
    ('COMERCIO_SERVICO', 'Sala comercial', ['SALA COMERCIAL', 'SALAS COMERCIAIS']),
    ('COMERCIO_VAREJO', 'Loja', ['LOJA']),
    ('COMERCIO_VAREJO', 'Distribuidora', ['DISTRIBUI']),
    ('COMERCIO_VAREJO', 'Comércio', ['COMERCIO', 'COMÉRCIO']),
]


# ╔═══════════════════════════════════════════════════════════════╗
# ║  FUNÇÕES DE PROCESSAMENTO                                    ║
# ╚═══════════════════════════════════════════════════════════════╝

def load_cnefe(path: str, municipio: int = None) -> pd.DataFrame:
    """Carrega CSV do CNEFE com tipagem correta.

    v4.3 (M4): CEP/COD_SETOR/NUM_ENDERECO/COD_UNICO_ENDERECO tipados como str
    na leitura — CEP int64 perdia zero à esquerda (ex.: 01001000/SP) e quebrava
    o check de 8 dígitos da confiabilidade; COD_UNICO como str atende C5 (merge).
    """
    from cnefe_download import SEP_CNEFE
    sep = SEP_CNEFE          # fonte unica: leitura e escrita nao podem divergir
    compression = 'gzip' if path.endswith('.gz') else None
    # v4.9.4 — a LISTA INVERTE-SE: é numérico o que está DECLARADO numérico;
    # todo o resto é TEXTO. A lista de exceções `_STR_COLS` protegia 4 colunas
    # e deixava `VAL_COMP_ELEM1..5` — o valor do apartamento, do bloco, do lote
    # — à inferência do pandas. Num arquivo em que esses valores são todos
    # numéricos e ALGUMA linha vem sem complemento, a coluna vira float64 e o
    # apartamento 101 passa a ser '101.0'. Como `num_unidade()` exige dígitos
    # puros, a grade para de ler aquele endereço e TODAS as hipóteses dele
    # desaparecem em silêncio. Encontrado com uma padaria ao lado de um prédio.
    #
    # Inverter a lista também é fail-closed para o futuro: coluna nova do IBGE
    # chega como texto, e texto se converte; float arredondado não se desfaz.
    _NUM_COLS = {'COD_UF', 'COD_MUNICIPIO', 'COD_ESPECIE', 'COD_TIPO_ESPECIE',
                 'COD_TIPO_ESPECI', 'COD_INDICADOR_ESTAB_ENDERECO',
                 'NV_GEO_COORD', 'LATITUDE', 'LONGITUDE'}
    def _limpo(c):
        return c.strip().replace('\r', '').strip('"')
    hdr = pd.read_csv(path, sep=sep, encoding='utf-8', nrows=0, compression=compression)
    dtypes = {c: str for c in hdr.columns if _limpo(c) not in _NUM_COLS}
    df = pd.read_csv(path, sep=sep, encoding='utf-8', low_memory=False,
                     compression=compression, dtype=dtypes)
    df.columns = [c.strip().replace('\r', '').strip('"') for c in df.columns]

    # Renomear coluna truncada se existir
    renames = {'COD_TIPO_ESPECI': 'COD_TIPO_ESPECIE', 'COD_TIPO_ESPEC': 'COD_TIPO_ESPECIE'}
    df.rename(columns={k: v for k, v in renames.items() if k in df.columns}, inplace=True)

    # ── R6: DEDUPE DA LINHA BRUTA, ANTES de qualquer coerção ──────────────
    # A doutrina dizia "dedup exato antes de qualquer estatística", mas o
    # dedupe rodava DEPOIS da normalização de NUM_ENDERECO. "001" e "1" são
    # duas linhas brutas DIFERENTES; viravam 1 e 1 e colapsavam em uma só.
    # Isso não é dedupe de linha bruta, é dedupe pós-transformação destrutiva
    # — e o que se perde nele é justamente o sinal de inconsistência da fonte.
    _n_raw = len(df)
    df = df.drop_duplicates().reset_index(drop=True)
    _n_dup_raw = _n_raw - len(df)

    # R1 na ingestão: NUM nulo/não-numérico -> 0, int garantido para TODOS os
    # fluxos (runner, coletivas standalone, mapa). Lido como str p/ não virar
    # float '104.0'; convertido aqui de forma controlada.
    #
    # NUM_ENDERECO_RAW preserva a grafia original: sem ele, "S/N", "ABC", "0"
    # e nulo viram o mesmo 0 e o pipeline perde a distinção entre SEM NÚMERO
    # DECLARADO, AUSENTE e MALFORMADO.
    if 'NUM_ENDERECO' in df.columns:
        df['NUM_ENDERECO_RAW'] = df['NUM_ENDERECO'].astype(str).str.strip()
        _num = pd.to_numeric(df['NUM_ENDERECO'], errors='coerce')
        df['NUM_STATUS'] = np.where(
            _num.notna() & (_num > 0), 'OK',
            np.where(df['NUM_ENDERECO_RAW'].str.upper().isin(
                         ['', 'S/N', 'SN', 'S.N.', '0', 'NAN', 'NONE']),
                     'SEM_NUMERO_DECLARADO', 'MALFORMADO'))
        df['NUM_ENDERECO'] = _num.fillna(0).astype(int)

    if municipio:
        df = df[pd.to_numeric(df['COD_MUNICIPIO'], errors='coerce') == int(municipio)].copy()
        if df.empty:
            print(f"ERRO: nenhum registro para COD_MUNICIPIO={municipio}")
            sys.exit(1)

    print(f"  Carregado: {len(df):,} registros")
    return df


REMERGE_RAIO_M = 50.0     # medido: 9,5% das colisoes de POA ficam abaixo disto


def _remerge_chave_por_proximidade(df, raio_m=REMERGE_RAIO_M):
    """Reune chaves que so diferem no TIPO/TITULO e coabitam no espaco.

    Sem isto, incluir o tipo na chave trocaria um erro por outro: separaria
    "RUA X 100" de "AV X 100" mesmo quando sao a mesma esquina grafada de dois
    jeitos pelo recenseador. A decisao e' da COORDENADA (R7), nunca do texto.
    """
    if 'CHAVE' not in df.columns or 'LATITUDE' not in df.columns:
        return df
    base = (df['COD_MUNICIPIO'].astype(str) + '|'
            + df['CHAVE'].str.split('|').str[3] + '|'
            + df['NUMERO'].astype(str) + '|'
            + df['CHAVE'].str.split('|').str[5])
    df = df.assign(_BASE=base)
    lat = pd.to_numeric(df['LATITUDE'], errors='coerce')
    lon = pd.to_numeric(df['LONGITUDE'], errors='coerce')
    # SÓ GEOMETRIA SANEADA DECIDE FUSÃO. A rotina aceitava qualquer par
    # lat/lon não nulo — inclusive coordenada que a própria `sanidade_geo`,
    # que roda ANTES no pipeline, já havia marcado como insana. A informação
    # existia e não estava sendo usada: um ponto reprovado ajudava a decidir
    # se duas chaves viravam uma. Se a coluna não existir (uso avulso da
    # função), o comportamento antigo permanece — declarado, não implícito.
    _ok = (df['SANIDADE_GEO'].astype(str).eq('OK')
           if 'SANIDADE_GEO' in df.columns
           else pd.Series(True, index=df.index))
    cen = (df.assign(_la=lat.where(_ok), _lo=lon.where(_ok))
             .dropna(subset=['_la', '_lo'])
             .groupby(['_BASE', 'CHAVE'])[['_la', '_lo']].mean().reset_index())
    nvar = cen.groupby('_BASE')['CHAVE'].transform('nunique')
    cand = cen[nvar > 1]
    if not len(cand):
        return df.drop(columns=['_BASE'])

    R = 6_371_000.0
    trocas = {}
    conta = df.groupby('CHAVE').size()
    for b, g in cand.groupby('_BASE'):
        la = np.radians(g['_la'].values); lo = np.radians(g['_lo'].values)
        dla = la[:, None] - la[None, :]; dlo = lo[:, None] - lo[None, :]
        a = (np.sin(dla / 2) ** 2
             + np.cos(la[:, None]) * np.cos(la[None, :]) * np.sin(dlo / 2) ** 2)
        dist = 2 * R * np.arcsin(np.sqrt(np.clip(a, 0, 1)))
        if np.nanmax(dist) > raio_m:
            continue                       # lugares distintos: seguem separados
        alvo = max(g['CHAVE'], key=lambda c: conta.get(c, 0))   # tipo modal vence
        for c in g['CHAVE']:
            if c != alvo:
                trocas[c] = alvo
    if trocas:
        df['CHAVE'] = df['CHAVE'].replace(trocas)
    return df.drop(columns=['_BASE'])


def normalizar_complemento(df: pd.DataFrame) -> pd.DataFrame:
    """Parseia os 5 pares NOM/VAL e reslota por camada semântica."""
    kws_by_layer = {}
    for kw, lay in LAYER_MAP.items():
        kws_by_layer.setdefault(lay, []).append(kw)

    for cam in CAMADAS:
        df[f'{cam}_TIPO'] = None
        df[f'{cam}_VALOR'] = None

    for i in range(1, 6):
        nom_upper = df[f'NOM_COMP_ELEM{i}'].fillna('').str.strip().str.upper()
        for cam in CAMADAS:
            mask = nom_upper.isin(kws_by_layer.get(cam, [])) & df[f'{cam}_TIPO'].isna()
            if mask.any():
                df.loc[mask, f'{cam}_TIPO'] = nom_upper[mask]
                df.loc[mask, f'{cam}_VALOR'] = (
                    df.loc[mask, f'VAL_COMP_ELEM{i}'].fillna('').astype(str).str.strip()
                )

    def _build(row):
        parts = []
        for c in CAMADAS:
            t, v = row[f'{c}_TIPO'], row[f'{c}_VALOR']
            if pd.notna(t) and t:
                parts.append(f"{t} {v}".strip() if v else t)
        return ' → '.join(parts)

    df['COMPLEMENTO_NORM'] = df.apply(_build, axis=1)
    return df


def classificar_uso(df: pd.DataFrame) -> pd.DataFrame:
    """Classifica macro-uso e atividade detalhada (classificação lexical de atividade)."""

    def _uso(row):
        esp = row['COD_ESPECIE']
        if esp == 1: return ('RESIDENCIAL', 'Domicílio particular')
        if esp == 2: return ('DOMICILIO_COLETIVO', 'Domicílio coletivo')
        if esp == 3: return ('AGROPECUARIO', 'Estab. agropecuário')
        if esp == 4: return ('EDUCACAO', 'Estab. de ensino')
        if esp == 5: return ('SAUDE', 'Estab. de saúde')
        if esp == 7: return ('CONSTRUCAO', 'Construção/reforma')
        if esp == 8: return ('RELIGIOSO', 'Estab. religioso')

        d = str(row.get('DSC_ESTABELECIMENTO', '')).upper().strip()
        if not d or d == 'NAN':
            # "Nao sei" nao pode virar "comercio". O CNEFE e' cadastro
            # ESTATISTICO de enderecos, nao cadastro comercial: espécie 6 sem
            # descricao diz que ha um estabelecimento, nao QUAL. Chamar isso de
            # COMERCIO_SERVICO fabricava falso positivo comercial em massa.
            return ('NAO_CLASSIFICADO', 'Sem descricao anotada')

        for rule in _LEXICO_ATIVIDADE:
            setor, ativ = rule[0], rule[1]
            kws = rule[2]
            match_type = rule[3] if len(rule) > 3 else 'contains'
            for kw in kws:
                if match_type == 'exact_or_prefix':
                    if d == kw or d.startswith(kw + ' '):
                        return (setor, ativ)
                else:
                    if kw in d:
                        if kw == 'LOJA' and 'VAGO' in d:
                            continue
                        return (setor, ativ)

        # Descricao existe mas nenhuma regra do lexico casou: a atividade e'
        # REAL e o setor e' DESCONHECIDO. Preserva-se o texto anotado, que e' a
        # evidencia; a classificacao fica em aberto ate haver regra.
        return ('NAO_CLASSIFICADO', d[:50])

    result = df.apply(_uso, axis=1, result_type='expand')
    df['SETOR_ATIVIDADE'] = result[0]
    df['ATIVIDADE_DETALHE'] = result[1]
    return df


def enriquecer(df: pd.DataFrame) -> pd.DataFrame:
    """Adiciona colunas descritivas e chaves de agrupamento."""
    df['QUALIDADE_COORD'] = df['NV_GEO_COORD'].map(QUALIDADE_GEO)
    df['DSC_TIPO_ESPECIE'] = df['COD_TIPO_ESPECIE'].map(TIPO_ESP)
    df['DSC_ESPECIE'] = df['COD_ESPECIE'].map(ESPECIE)
    df['LOGRADOURO'] = (
        df['NOM_TIPO_SEGLOGR'].fillna('') + ' ' +
        df['NOM_TITULO_SEGLOGR'].fillna('') + ' ' +
        df['NOM_SEGLOGR'].fillna('')
    ).str.strip()
    df['NUMERO'] = pd.to_numeric(df['NUM_ENDERECO'], errors='coerce').fillna(0).astype(int)
    # v4.3: CHAVE usa o logradouro HARMONIZADO quando disponível (typo não
    # fragmenta mais o cluster). CEP fica FORA do escopo decisório da chave —
    # é entrada digitada; a coordenada é quem confirma (doutrina R7).
    # N14: prefixo de MUNICÍPIO — sem ele, RUA A|100|CENTRO de duas cidades
    # colidiria num run multi-município.
    _logr_key = (df['NOM_SEGLOGR_HARM'] if 'NOM_SEGLOGR_HARM' in df.columns
                 else df['NOM_SEGLOGR'])
    _mun_key = (pd.to_numeric(df['COD_MUNICIPIO'], errors='coerce')
                .fillna(0).astype('int64').astype(str)
                if 'COD_MUNICIPIO' in df.columns else '0')
    # TIPO e TÍTULO entram na chave (v4.8). Sem eles, RUA PRESIDENTE VARGAS 100
    # e AVENIDA PRESIDENTE VARGAS 100 caíam no mesmo grupo. Medido em Porto
    # Alegre: 714 chaves fundindo tipos distintos, 3.001 endereços, mediana de
    # 687 m entre os blocos fundidos e 80,6% acima de 200 m — colisão real, não
    # variação de grafia. E o dano não para na identificação: QTD_END, blocos,
    # uso misto, coletiva, S2/S3, declarado×contado, gaps, porte e RCC inteiro
    # rodam por CHAVE.
    _tipo = df['NOM_TIPO_SEGLOGR'].fillna('').astype(str).map(strip_accents).str.upper().str.strip()
    _tit = df.get('NOM_TITULO_SEGLOGR', pd.Series('', index=df.index)) \
             .fillna('').astype(str).map(strip_accents).str.upper().str.strip()
    df['CHAVE'] = (
        _mun_key + '|' + _tipo + '|' + _tit + '|' +
        _logr_key.fillna('').astype(str).str.upper().str.strip() + '|' +
        df['NUMERO'].astype(str) + '|' +
        df['DSC_LOCALIDADE'].fillna('').str.upper().str.strip()
    )

    # Re-merge por proximidade. Acrescentar o tipo cru QUEBRARIA os grupos
    # legítimos em que o recenseador escreveu RUA numa linha e AV na outra para
    # o MESMO lugar — 9,5% das colisões de POA têm centróides a menos de 50 m.
    # A geometria arbitra, como manda R7: mesmo nome, mesmo número, mesma
    # localidade e centróides próximos voltam a ser um só grupo, sob o tipo
    # modal. Longe, permanecem separados.
    df = _remerge_chave_por_proximidade(df)
    # QTD_END calculado apenas entre registros com NUMERO > 0
    # NUMERO=0 são endereços sem numeração oficial que inflam agrupamentos
    df['QTD_END'] = 1  # default
    mask_valid = df['NUMERO'] > 0
    df.loc[mask_valid, 'QTD_END'] = (
        df[mask_valid].groupby('CHAVE')['CHAVE'].transform('count')
    )
    df['N_BLOCOS'] = 0
    df.loc[mask_valid, 'N_BLOCOS'] = (
        df[mask_valid].groupby('CHAVE')['BLOCO_VALOR'].transform(
            lambda x: x.replace('', np.nan).dropna().nunique()
        )
    )
    df['FLAG_MB'] = (df['N_BLOCOS'] >= 2).astype(int)
    return df


def classificar_coletiva(df: pd.DataFrame) -> pd.DataFrame:
    """Classifica cada registro como coletiva com tipo e confiança."""

    def _clf(r):
        q, qu, te, nv, mb = r['QTD_END'], r['QUALIDADE_COORD'], r.get('COD_TIPO_ESPECIE'), r['NV_GEO_COORD'], r['FLAG_MB']
        has = {c: pd.notna(r.get(f'{c}_TIPO')) for c in CAMADAS}
        ind = r.get('COD_INDICADOR_ESTAB_ENDERECO')

        if q <= 1:
            if te == 103 and qu == 'VALIDADA':
                return ('APTO_ISOLADO', 'MEDIA', 'tipo=103 único')
            if pd.notna(ind) and ind in (2, 3, 4) and qu == 'VALIDADA':
                return ('MULT_ESTAB', 'ALTA', f'ind={int(ind)}')
            return ('NAO_COLETIVA', 'NAO_COLETIVA', '')

        # ── ALTA
        if mb and qu == 'VALIDADA':
            n = r['N_BLOCOS']
            if te == 103: return ('COND_VERT_MULTIBLOCOS', 'ALTA', f'{n} blocos+apto')
            if has['MORADIA'] or has['VIA_INTERNA']: return ('COND_HORIZ_MULTIBLOCOS', 'ALTA', f'{n} blocos/vias')
            return ('COND_MULTIBLOCOS', 'ALTA', f'{n} blocos')
        if has['AGRUPAMENTO'] and (has['BLOCO'] or has['UNIDADE']) and qu == 'VALIDADA':
            return ('EMPREEND_NOMEADO', 'ALTA', f'{r["AGRUPAMENTO_TIPO"]}+blocos')
        if te == 103 and qu == 'VALIDADA': return ('VERT_APTO', 'ALTA', 'tipo=103+validada')
        if nv == 2: return ('VERT_APTO', 'ALTA', 'NV_GEO=2')
        if te == 102 and qu == 'VALIDADA': return ('HORIZ_VILA_COND', 'ALTA', 'tipo=102+validada')
        if has['UNIDADE'] and qu == 'VALIDADA': return ('VERT_COMPL', 'ALTA', f'unid={r["UNIDADE_TIPO"]}')
        if has['BLOCO'] and qu == 'VALIDADA': return ('VERT_BLOCO', 'ALTA', 'bloco+validada')
        if pd.notna(ind) and ind in (2, 3, 4) and qu == 'VALIDADA': return ('MULT_ESTAB', 'ALTA', f'ind={int(ind)}')
        if has['VIA_INTERNA'] and qu == 'VALIDADA': return ('COND_VIA_INT', 'ALTA', 'via interna')

        # ── MÉDIA
        if has['MORADIA'] and qu == 'VALIDADA': return ('HORIZ_MORADIAS', 'MEDIA', r['MORADIA_TIPO'])
        if has['POSICAO'] and qu == 'VALIDADA': return ('HORIZ_SUBDIV', 'MEDIA', r['POSICAO_TIPO'])
        if te == 103 and qu == 'ESTIMADA': return ('VERT_APTO', 'MEDIA', 'tipo=103+ESTIMADA')
        if mb and qu in ('ESTIMADA', 'BAIXA'): return ('COND_MULTIBLOCOS', 'MEDIA', f'{r["N_BLOCOS"]}bl+{qu}')
        if q > 2 and qu == 'VALIDADA': return ('AGRUP_END', 'MEDIA', f'N={q} s/compl')

        # v4.3: DUPLA_END refinado — 2 registros com split de pavimento
        # (TERREO/ANDAR) e coordenada validada sobem para MEDIA.
        # POSICAO+VALIDADA já sobe antes via HORIZ_SUBDIV.
        if q == 2 and has['PAVIMENTO'] and qu == 'VALIDADA':
            return ('DUPLA_END', 'MEDIA', 'pavimento split')

        # ── BAIXA
        if qu in ('ESTIMADA', 'BAIXA'): return ('AGRUP_TEXTO', 'BAIXA', f'N={q}+{qu}')
        if q == 2: return ('DUPLA_END', 'BAIXA', '2 reg')

        return ('NAO_COLETIVA', 'NAO_COLETIVA', '')

    result = df.apply(_clf, axis=1, result_type='expand')
    df['TIPO_COLETIVA'] = result[0]
    df['CONFIANCA'] = result[1]
    df['EVIDENCIA'] = result[2]
    df['FLAG_COL'] = (df['CONFIANCA'] != 'NAO_COLETIVA').astype(int)
    df['TIPO_COLETIVA_DESC'] = df['TIPO_COLETIVA'].map(TIPO_DESC)
    df['CONFIANCA_DESC'] = df['CONFIANCA'].map(CONF_DESC)

    # ID_COLETIVA
    chaves_unicas = df[df['FLAG_COL'] == 1]['CHAVE'].drop_duplicates().reset_index(drop=True)
    id_map = {ch: i + 1 for i, ch in enumerate(chaves_unicas)}
    df['ID_COLETIVA'] = df['CHAVE'].map(id_map)

    # ── NUMERO=0: forçar como não-coletiva (sem numeração, agrupamento falso)
    mask_zero = df['NUMERO'] == 0
    df.loc[mask_zero, 'FLAG_COL'] = 0
    df.loc[mask_zero, 'TIPO_COLETIVA'] = 'NAO_COLETIVA'
    df.loc[mask_zero, 'CONFIANCA'] = 'NAO_COLETIVA'
    df.loc[mask_zero, 'TIPO_COLETIVA_DESC'] = TIPO_DESC['NAO_COLETIVA']
    df.loc[mask_zero, 'CONFIANCA_DESC'] = CONF_DESC['NAO_COLETIVA']
    df.loc[mask_zero, 'ID_COLETIVA'] = np.nan

    # Reclassificar registros que ficaram sozinhos após exclusão
    mask_solo = (df['QTD_END'] == 1) & (df['FLAG_COL'] == 1) & (df['TIPO_COLETIVA'] != 'APTO_ISOLADO') & (df['TIPO_COLETIVA'] != 'MULT_ESTAB')
    df.loc[mask_solo, 'FLAG_COL'] = 0
    df.loc[mask_solo, 'TIPO_COLETIVA'] = 'NAO_COLETIVA'
    df.loc[mask_solo, 'CONFIANCA'] = 'NAO_COLETIVA'
    df.loc[mask_solo, 'TIPO_COLETIVA_DESC'] = TIPO_DESC['NAO_COLETIVA']
    df.loc[mask_solo, 'CONFIANCA_DESC'] = CONF_DESC['NAO_COLETIVA']
    df.loc[mask_solo, 'ID_COLETIVA'] = np.nan

    # Regenerar ID_COLETIVA limpo
    # R8: Int64 nullable, NUNCA float. O `.map` com chave ausente devolveria
    # NaN e tiparia a coluna como float64 — aqui os valores são pequenos e
    # não perderiam dígito, mas a REGRA é de representação, não de magnitude:
    # id em float é o vetor do defeito, e abrir exceção por tamanho é como a
    # classe reaparece. Efeito colateral bom: some o '2.0' do Excel.
    chaves_final = df[df['FLAG_COL'] == 1]['CHAVE'].drop_duplicates().reset_index(drop=True)
    id_map2 = {ch: i + 1 for i, ch in enumerate(chaves_final)}
    df['ID_COLETIVA'] = df['CHAVE'].map(id_map2).astype('Int64')

    return df


# ╔═══════════════════════════════════════════════════════════════╗
# ║  AGRUPAMENTO                                                  ║
# ╚═══════════════════════════════════════════════════════════════╝

def agrupar_enderecos(df: pd.DataFrame) -> pd.DataFrame:
    """Agrupa coletivas: 1 linha por bloco com centróide e contagem por uso."""
    df_col = df[(df['FLAG_COL'] == 1) & (df['NUMERO'] > 0)].copy()
    df_col['BLOCO_KEY'] = df_col['BLOCO_VALOR'].fillna('').replace('', np.nan).fillna('-')

    _aggs_extra = {}
    if 'RADAR_ID_ENDERECO' in df_col.columns:
        _aggs_extra['RADAR_ID_ENDERECO'] = ('RADAR_ID_ENDERECO', 'first')
    if 'RADAR_ID_BLOCO' in df_col.columns:
        _aggs_extra['RADAR_ID_BLOCO'] = ('RADAR_ID_BLOCO', 'first')
    if 'COD_INDICADOR_ESTAB_ENDERECO' in df_col.columns:
        # N16: pior caso declarado de multiplicidade de estabelecimentos no bloco
        _aggs_extra['IND_ESTAB_MAX'] = (
            'COD_INDICADOR_ESTAB_ENDERECO',
            lambda x: pd.to_numeric(x, errors='coerce').max())
    if 'VALIDACAO_ESTAB_DECL' in df_col.columns:
        # N16: veredito declarado×contado (constante por CHAVE)
        _aggs_extra['VALIDACAO_ESTAB_DECL'] = ('VALIDACAO_ESTAB_DECL', 'first')
    grp = df_col.groupby(['CHAVE', 'BLOCO_KEY']).agg(
        **_aggs_extra,
        ID_COLETIVA=('ID_COLETIVA', 'first'),
        LOGRADOURO=('LOGRADOURO', 'first'),
        NUMERO=('NUMERO', 'first'),
        LOCALIDADE=('DSC_LOCALIDADE', 'first'),
        CEP=('CEP', 'first'),
        SETOR=('COD_SETOR', 'first'),
        QTD_REGISTROS=('COD_UNICO_ENDERECO', 'count'),
        TIPO_COLETIVA=('TIPO_COLETIVA_DESC', 'first'),
        CONFIANCA=('CONFIANCA_DESC', 'first'),
        CONFIANCA_COD=('CONFIANCA', 'first'),
        QUALIDADE_COORD=('QUALIDADE_COORD', 'first'),
        BLOCO_TIPO=('BLOCO_TIPO', lambda x: next((str(v) for v in x.dropna() if str(v).strip()), '-')),
        N_ANDARES=('PAVIMENTO_VALOR', lambda x: x.replace('', np.nan).dropna().nunique()),
        N_UNIDADES=('UNIDADE_VALOR', lambda x: x.replace('', np.nan).dropna().nunique()),
        N_MORADIAS=('MORADIA_VALOR', lambda x: x.replace('', np.nan).dropna().nunique()),
        QTD_RESIDENCIAL=('SETOR_ATIVIDADE', lambda x: (x == 'RESIDENCIAL').sum()),
        QTD_COMERCIAL=('SETOR_ATIVIDADE', lambda x: x.isin(['COMERCIO_VAREJO', 'COMERCIO_SERVICO', 'BELEZA_ESTETICA']).sum()),
        QTD_INDUSTRIAL=('SETOR_ATIVIDADE', lambda x: (x == 'INDUSTRIAL').sum()),
        QTD_ALIMENTACAO=('SETOR_ATIVIDADE', lambda x: (x == 'ALIMENTACAO').sum()),
        QTD_SERVICO=('SETOR_ATIVIDADE', lambda x: x.isin(['SERVICO_PROF', 'SERVICO_AUTO']).sum()),
        QTD_MANUFATURA=('SETOR_ATIVIDADE', lambda x: (x == 'MANUFATURA').sum()),
        QTD_SAUDE=('SETOR_ATIVIDADE', lambda x: (x == 'SAUDE').sum()),
        QTD_EDUCACAO=('SETOR_ATIVIDADE', lambda x: (x == 'EDUCACAO').sum()),
        QTD_RELIGIOSO=('SETOR_ATIVIDADE', lambda x: (x == 'RELIGIOSO').sum()),
        QTD_HOSPEDAGEM=('SETOR_ATIVIDADE', lambda x: (x == 'HOSPEDAGEM').sum()),
        QTD_VAGO=('SETOR_ATIVIDADE', lambda x: (x == 'VAGO').sum()),
        QTD_OUTROS=('SETOR_ATIVIDADE', lambda x: x.isin(['LAZER', 'ASSOCIACAO_ONG', 'DIVERSOS', 'AGROPECUARIO', 'CONSTRUCAO', 'DOMICILIO_COLETIVO']).sum()),
        CENTROIDE_LAT=('LATITUDE', 'mean'),
        CENTROIDE_LON=('LONGITUDE', 'mean'),
    ).reset_index()

    grp.rename(columns={'BLOCO_KEY': 'BLOCO'}, inplace=True)
    totais = grp.groupby('ID_COLETIVA').agg(
        TOTAL_GERAL=('QTD_REGISTROS', 'sum'),
        N_BLOCOS=('BLOCO', lambda x: (x != '-').sum()),
    ).reset_index()
    grp = grp.merge(totais, on='ID_COLETIVA', how='left')
    grp = grp.drop(columns=['CHAVE']).sort_values(['ID_COLETIVA', 'BLOCO'])

    col_order = [
        'ID_COLETIVA', 'LOGRADOURO', 'NUMERO', 'LOCALIDADE', 'CEP', 'SETOR',
        'BLOCO_TIPO', 'BLOCO', 'N_BLOCOS', 'QTD_REGISTROS', 'TOTAL_GERAL',
        'N_ANDARES', 'N_UNIDADES', 'N_MORADIAS',
        'QTD_RESIDENCIAL', 'QTD_COMERCIAL', 'QTD_INDUSTRIAL', 'QTD_ALIMENTACAO',
        'QTD_SERVICO', 'QTD_MANUFATURA', 'QTD_SAUDE', 'QTD_EDUCACAO',
        'QTD_RELIGIOSO', 'QTD_HOSPEDAGEM', 'QTD_VAGO', 'QTD_OUTROS',
        'TIPO_COLETIVA', 'CONFIANCA', 'CONFIANCA_COD', 'QUALIDADE_COORD',
        'CENTROIDE_LAT', 'CENTROIDE_LON',
    ]
    col_order += [c for c in ('RADAR_ID_ENDERECO', 'RADAR_ID_BLOCO',
                              'IND_ESTAB_MAX', 'VALIDACAO_ESTAB_DECL')
                  if c in grp.columns]
    return grp[col_order]


# ╔═══════════════════════════════════════════════════════════════╗
# ║  GERAÇÃO DO EXCEL FORMATADO                                  ║
# ╚═══════════════════════════════════════════════════════════════╝

# Paleta de estilos
_HF = Font(name='Arial', bold=True, size=10, color='FFFFFF')
_HFILL = PatternFill('solid', fgColor='2F5496')
_TF = Font(name='Arial', bold=True, size=13, color='2F5496')
_SF = Font(name='Arial', size=9, color='595959')
_BF = Font(name='Arial', size=9)
_SECF = Font(name='Arial', bold=True, size=11, color='2F5496')
_BORDER = Border(bottom=Side(style='thin', color='E0E0E0'))

_CONF_FILLS = {
    'ALTA': PatternFill('solid', fgColor='C6EFCE'),
    'MEDIA': PatternFill('solid', fgColor='FFEB9C'),
    'BAIXA': PatternFill('solid', fgColor='FFC7CE'),
}
_SETOR_FILLS = {
    'ALIMENTACAO': PatternFill('solid', fgColor='FCE5CD'),
    'INDUSTRIAL': PatternFill('solid', fgColor='F4CCCC'),
    'COMERCIO_VAREJO': PatternFill('solid', fgColor='FFF2CC'),
    'COMERCIO_SERVICO': PatternFill('solid', fgColor='FFF2CC'),
    'SERVICO_PROF': PatternFill('solid', fgColor='D9EAD3'),
    'SERVICO_AUTO': PatternFill('solid', fgColor='D9EAD3'),
    'MANUFATURA': PatternFill('solid', fgColor='EAD1DC'),
    'BELEZA_ESTETICA': PatternFill('solid', fgColor='EAD1DC'),
    'SAUDE': PatternFill('solid', fgColor='CFE2F3'),
    'EDUCACAO': PatternFill('solid', fgColor='CFE2F3'),
    'RELIGIOSO': PatternFill('solid', fgColor='D9D2E9'),
    'HOSPEDAGEM': PatternFill('solid', fgColor='FCE5CD'),
    'VAGO': PatternFill('solid', fgColor='E0E0E0'),
    'AGROPECUARIO': PatternFill('solid', fgColor='D9EAD3'),
}


def _write_sheet(wb, name, title, subtitle, df_data, start_row=3,
                 conf_col=None, setor_col=None):
    """Escreve uma aba com dados e formatação profissional."""
    ws = wb.create_sheet(name)
    ws.cell(1, 1, title).font = _TF
    ws.cell(2, 1, subtitle).font = _SF

    cols = list(df_data.columns)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(start_row, j, c)
        cell.font, cell.fill = _HF, _HFILL
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    conf_idx = cols.index(conf_col) + 1 if conf_col and conf_col in cols else None
    setor_idx = cols.index(setor_col) + 1 if setor_col and setor_col in cols else None

    for i, tup in enumerate(df_data.itertuples(index=False), start_row + 1):
        for j, val in enumerate(tup, 1):
            if isinstance(val, (np.integer, np.int64)):
                val = int(val)
            elif isinstance(val, (np.floating, np.float64)):
                val = round(float(val), 7) if not np.isnan(val) else None
            elif isinstance(val, (np.bool_,)):
                val = int(val)
            elif pd.isna(val):
                val = None
            cell = ws.cell(i, j, val)
            cell.border = _BORDER
            if isinstance(val, (int, float)) and val is not None:
                cell.number_format = '#,##0'

        if conf_idx:
            v = str(ws.cell(i, conf_idx).value or '')
            for key, fill in _CONF_FILLS.items():
                if key == v or key.capitalize() in v or f'— {key.capitalize()}' in v.replace('á', 'a').replace('é', 'e'):
                    ws.cell(i, conf_idx).fill = fill
                    break
        if setor_idx:
            v = str(ws.cell(i, setor_idx).value or '')
            if v in _SETOR_FILLS:
                ws.cell(i, setor_idx).fill = _SETOR_FILLS[v]

    ws.freeze_panes = ws.cell(start_row + 1, 1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{start_row + len(df_data)}"

    # Auto-width
    for col_cells in ws.columns:
        lens = [len(str(c.value or '')) for c in col_cells[:200]]
        if lens:
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(lens) + 2, 40)

    return ws


def gerar_excel(df: pd.DataFrame, grp: pd.DataFrame, output_path: str):
    """Gera Excel final com 4 abas formatadas."""
    wb = Workbook()

    # ── ABA RESUMO ──
    ws1 = wb.active
    ws1.title = 'Resumo'
    cod_mun = df['COD_MUNICIPIO'].iloc[0]
    n_col = df['FLAG_COL'].sum()
    n_end = grp['ID_COLETIVA'].nunique()
    n_nres = len(df[(df['SETOR_ATIVIDADE'] != 'RESIDENCIAL') & (df['SETOR_ATIVIDADE'] != 'CONSTRUCAO')])

    ws1.cell(1, 1, f'Radar Coletivo — Município {cod_mun} — Ligações Coletivas').font = _TF
    ws1.cell(2, 1, f'{len(df):,} registros | {n_col:,} coletivos | {n_end:,} endereços | {n_nres:,} não-residencial').font = _SF

    r = 4
    ws1.cell(r, 1, 'NORMALIZAÇÃO — 7 CAMADAS SEMÂNTICAS').font = _SECF
    r += 1
    for cam, desc in [
        ('AGRUPAMENTO', 'Empreendimento → PREDIO, EDIFICIO, QUADRA, PAVILHAO'),
        ('VIA_INTERNA', 'Via interna → ALAMEDA INTERNA, RUA INTERNA'),
        ('BLOCO', 'Subdivisão → BLOCO, TORRE, ENTRADA'),
        ('PAVIMENTO', 'Vertical → ANDAR, TERREO, SUBSOLO, COBERTURA'),
        ('UNIDADE', 'Economia → APARTAMENTO, SALA, LOJA, QUITINETE, BOX'),
        ('MORADIA', 'Horizontal → CASA, LOTE, SOBRADO, CHALE'),
        ('POSICAO', 'Posição → FUNDOS, FRENTE, LADO, ANEXO'),
    ]:
        ws1.cell(r, 1, cam).font = _BF
        ws1.cell(r, 2, desc).font = _BF
        r += 1

    r += 1
    ws1.cell(r, 1, 'CLASSIFICAÇÃO').font = _SECF
    r += 1
    ws1.cell(r, 1, 'Tipo')
    ws1.cell(r, 2, 'Confiança')
    ws1.cell(r, 3, 'Qtd')
    for c in range(1, 4):
        ws1.cell(r, c).font = _HF
        ws1.cell(r, c).fill = _HFILL
    r += 1
    for _, rd in df.groupby(['TIPO_COLETIVA_DESC', 'CONFIANCA_DESC']).size().reset_index(name='N').sort_values('N', ascending=False).iterrows():
        ws1.cell(r, 1, rd['TIPO_COLETIVA_DESC']).font = _BF
        ws1.cell(r, 2, rd['CONFIANCA_DESC']).font = _BF
        ws1.cell(r, 3, int(rd['N'])).number_format = '#,##0'
        r += 1

    r += 1
    ws1.cell(r, 1, 'SETORES NÃO-RESIDENCIAL').font = _SECF
    r += 1
    df_nres_tmp = df[(df['SETOR_ATIVIDADE'] != 'RESIDENCIAL') & (df['SETOR_ATIVIDADE'] != 'CONSTRUCAO')]
    if 'RADAR_CONF_FAIXA' in df_nres_tmp.columns:   # mesmo critério F7 da aba
        df_nres_tmp = df_nres_tmp[df_nres_tmp['RADAR_CONF_FAIXA'].isin(['ALTA', 'MUITO_ALTA'])]
    for setor, n in df_nres_tmp['SETOR_ATIVIDADE'].value_counts().items():
        c1 = ws1.cell(r, 1, setor)
        c1.fill = _SETOR_FILLS.get(setor, PatternFill())
        ws1.cell(r, 2, int(n)).number_format = '#,##0'
        r += 1

    ws1.column_dimensions['A'].width = 55
    ws1.column_dimensions['B'].width = 55
    ws1.column_dimensions['C'].width = 12

    # ── ABA ENDERECOS COLETIVOS ──
    _write_sheet(
        wb, 'Enderecos_Coletivos',
        f'ENDEREÇOS COLETIVOS — 1 LINHA POR BLOCO — {len(grp):,} linhas',
        'ID_COLETIVA agrupa. Cada bloco tem centróide e contagem por setor.',
        grp, conf_col='CONFIANCA_COD',
    )

    # ── ABA COLETIVAS UNIDADES ──
    df_col = df[df['FLAG_COL'] == 1].copy()
    col_cols = [
        'ID_COLETIVA', 'COD_UNICO_ENDERECO', 'DSC_LOCALIDADE', 'LOGRADOURO', 'NUMERO',
        'CEP', 'COD_SETOR',
        'AGRUPAMENTO_TIPO', 'AGRUPAMENTO_VALOR', 'VIA_INTERNA_TIPO', 'VIA_INTERNA_VALOR',
        'BLOCO_TIPO', 'BLOCO_VALOR', 'PAVIMENTO_TIPO', 'PAVIMENTO_VALOR',
        'UNIDADE_TIPO', 'UNIDADE_VALOR', 'MORADIA_TIPO', 'MORADIA_VALOR',
        'POSICAO_TIPO', 'POSICAO_VALOR', 'COMPLEMENTO_NORM',
        'LATITUDE', 'LONGITUDE', 'NV_GEO_COORD', 'QUALIDADE_COORD',
        'SETOR_ATIVIDADE', 'ATIVIDADE_DETALHE', 'DSC_ESTABELECIMENTO', 'DSC_TIPO_ESPECIE',
        'QTD_END', 'N_BLOCOS', 'FLAG_MB', 'TIPO_COLETIVA_DESC', 'CONFIANCA_DESC',
    ]
    df_col_out = df_col[col_cols].sort_values(
        ['CONFIANCA_DESC', 'ID_COLETIVA', 'BLOCO_VALOR', 'PAVIMENTO_VALOR', 'UNIDADE_VALOR']
    )

    _write_sheet(
        wb, 'Coletivas_Unidades',
        f'CADA ECONOMIA COLETIVA — {len(df_col_out):,} registros',
        '7 camadas normalizadas + classificação lexical de atividade',
        df_col_out, conf_col='CONFIANCA_DESC', setor_col='SETOR_ATIVIDADE',
    )

    # ── ABA NÃO RESIDENCIAL ──
    df_nres = df[(df['SETOR_ATIVIDADE'] != 'RESIDENCIAL') & (df['SETOR_ATIVIDADE'] != 'CONSTRUCAO')].copy()
    # v4.5.2: mesmo critério F7 do runner — faixa de localização ALTA+
    if 'RADAR_CONF_FAIXA' in df_nres.columns:
        df_nres = df_nres[df_nres['RADAR_CONF_FAIXA'].isin(['ALTA', 'MUITO_ALTA'])]
    nres_cols = [
        'COD_UNICO_ENDERECO', 'DSC_LOCALIDADE', 'LOGRADOURO', 'NUMERO',
        'COMPLEMENTO_NORM', 'CEP', 'COD_SETOR',
        'LATITUDE', 'LONGITUDE', 'NV_GEO_COORD', 'QUALIDADE_COORD',
        'SETOR_ATIVIDADE', 'ATIVIDADE_DETALHE', 'DSC_ESTABELECIMENTO',
        'DSC_ESPECIE', 'DSC_TIPO_ESPECIE',
        'QTD_END', 'FLAG_COL', 'TIPO_COLETIVA_DESC', 'CONFIANCA_DESC',
    ]
    df_nres_out = df_nres[nres_cols].sort_values(['SETOR_ATIVIDADE', 'ATIVIDADE_DETALHE', 'DSC_LOCALIDADE'])

    _write_sheet(
        wb, 'Nao_Residencial',
        f'NÃO-RESIDENCIAL — {len(df_nres_out):,} registros',
        'SETOR_ATIVIDADE (macro) + ATIVIDADE_DETALHE (granular)',
        df_nres_out, conf_col='CONFIANCA_DESC', setor_col='SETOR_ATIVIDADE',
    )

    wb.save(output_path)
    print(f"\n✓ Excel salvo: {output_path}")
    print(f"  Abas: {wb.sheetnames}")


# ╔═══════════════════════════════════════════════════════════════╗
# ║  MAIN                                                         ║
# ╚═══════════════════════════════════════════════════════════════╝

def main():
    parser = argparse.ArgumentParser(description='Radar Coletivo — Análise de Ligações Coletivas')
    parser.add_argument('--input', required=True, help='CSV bruto (ou .csv.gz do download)')
    parser.add_argument('--municipio', type=int, default=None, help='Filtrar por código IBGE do município')
    parser.add_argument('--output', default=None, help='Caminho do Excel de saída')
    parser.add_argument('--pickle', default=None, help='Exportar pickle para uso no cnefe_mapa.py')
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"ERRO: arquivo não encontrado: {input_path}")
        sys.exit(1)

    # Determinar nome do output
    if args.output:
        output_path = args.output
    else:
        stem = input_path.stem.replace('.csv', '')
        output_path = str(input_path.parent / f'RADAR_{stem}_COLETIVAS.xlsx')  # R5

    print(f"Radar Coletivo — Coletivas — Análise de {input_path.name}")
    print(f"{'=' * 60}")

    # v4.5.2: este entrypoint é WRITER NÃO-CANÔNICO — consome o MESMO
    # núcleo preparar_base do runner (dedup→sanidade→anomalia→F7→
    # classificação→gate em um ÚNICO lugar; nada é recalculado aqui).
    import radar_pipeline as rp
    print("\n[1/3] Preparando base (núcleo único radar_pipeline)...")
    df, stats = rp.preparar_base(str(input_path), municipio=args.municipio,
                                 log=lambda m: print('      ' + m))
    m_gate = rp.gate_universal(df)
    print(f"      gate universal: {int(m_gate.sum()):,} / "
          f"{stats['apos_dedup']:,} pós-dedup ({stats['entrada_raw']:,} raw)")
    df = df[m_gate].copy()

    print("\n[2/3] Resumo da classificação:")
    summary = df.groupby(['CONFIANCA', 'TIPO_COLETIVA']).size().reset_index(name='N')
    print(summary.to_string(index=False))

    print("\n[3/3] Gerando Excel (writer legado/diagnóstico)...")
    grp = agrupar_enderecos(df)
    gerar_excel(df, grp, output_path)

    # Exportar pickle se solicitado
    if args.pickle:
        import pickle
        pickle.dump(df, open(args.pickle, 'wb'))
        print(f"  Pickle: {args.pickle}")

    print(f"\n{'=' * 60}")
    print(f"Concluído. {len(df):,} registros processados.")


if __name__ == '__main__':
    main()
