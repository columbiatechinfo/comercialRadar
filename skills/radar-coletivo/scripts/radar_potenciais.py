#!/usr/bin/env python3
"""
radar_potenciais.py — Runner canônico do Radar Coletivo (6 abas + manifest)
===========================================================================
Motor de DESCOBERTA DE CANDIDATOS sobre bases de endereços/domicílios
(município, UF ou multi-UF). A preparação da base vive em UM único lugar
(radar_pipeline.preparar_base): load tipado -> schema -> DEDUP exato ->
harmonização de grafia -> sanidade geo 3 camadas -> complementos 7 camadas
-> promoção Q/L -> classificação lexical de atividade -> chaves/RADAR_ID ->
confronto declarado×contado -> anomalia de coleta S1-S5 -> F7 confiabilidade
-> classificação de coletivas -> gate universal.

Este runner só orquestra os PRODUTOS: condomínio horizontal, agrupamento por
bloco, polos (DBSCAN por densidade de ENDEREÇOS), simples + hipóteses de
expansão (separadas do observado), gate_pos_condicoes (RadarQualityError,
imune a python -O) e o Excel de 6 abas + <saida>.manifest.json com linhagem,
funil por causa prioritária e validação runtime×lock.

Uso: python radar_potenciais.py --input <base.csv> --output <potenciais.xlsx>
     [--municipio COD] [--sanidade-pct 0.001] [--anomalia-seq-min 15]
     [--anomalia-cobertura 0.95] [--dbscan-eps 100] [--dbscan-min 5]
     [--geo-raio-grosseiro 300] [--source-dataset X --source-vintage Y]
"""
import sys, os, re, time, argparse
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment

import radar_utils as xu
import radar_pipeline as rp
import cnefe_coletivas as cc

t0 = time.time()
def log(msg): print(f"[{time.time()-t0:6.1f}s] {msg}", flush=True)

ap = argparse.ArgumentParser(
    description='Radar Coletivo — runner canônico (6 abas + manifest de linhagem)')
ap.add_argument('--input', required=True)
ap.add_argument('--output', required=True)
ap.add_argument('--municipio', type=int, default=None,
                help='Filtro opcional: processa só 1 código de município')
ap.add_argument('--sanidade-pct', type=float, default=0.001,
                help='Percentil da FLAG_OUTLIER_GEO informativa (por partição)')
ap.add_argument('--anomalia-seq-min', '--fraude-seq-min', dest='fraude_seq_min',
                type=int, default=15,
                help='S2: mínimo de UNIDADES DISTINTAS em sequência perfeita')
ap.add_argument('--anomalia-cobertura', '--fraude-cobertura', dest='fraude_cobertura',
                type=float, default=0.95,
                help='S2: cobertura mínima da sequência')
ap.add_argument('--dbscan-eps', type=float, default=100, help='T5: raio em metros')
ap.add_argument('--dbscan-min', type=int, default=5, help='T5: mínimo de endereços por polo')
ap.add_argument('--dbscan-peso-max', type=int, default=500,
                help='sem efeito desde v4.5 (correção de semântica); aceito por compat')
ap.add_argument('--geo-raio-grosseiro', type=float, default=300.0,
                help='T0: km à mediana da partição p/ ERRO grosseiro (dívida '
                     'aberta: substituir por malha municipal ST_Covers)')
ap.add_argument('--source-dataset', default=None,
                help='Declaração da fonte p/ o manifest (senão: inferência por '
                     'nome de arquivo ou UNKNOWN — nunca inventada)')
ap.add_argument('--source-vintage', default=None,
                help='Safra/vintage declarada da fonte p/ o manifest')
args = ap.parse_args()

# ════════════════════════════════════════════════════════════════
#  NÚCLEO ÚNICO (v4.5.2): preparar_base concentra load→schema→dedup→
#  harmonização→sanidade 3 camadas→complementos→lexical→chaves→RADAR_ID→
#  declarado×contado→anomalia→F7→classificação→gate. Os três entrypoints
#  (runner, coletivas, mapa) consomem ESTE mesmo núcleo — writers não
#  recalculam regra de negócio.
# ════════════════════════════════════════════════════════════════
df, stats = rp.preparar_base(
    args.input, municipio=args.municipio, sanidade_pct=args.sanidade_pct,
    anomalia_seq_min=args.fraude_seq_min, anomalia_cobertura=args.fraude_cobertura,
    geo_raio_grosseiro=args.geo_raio_grosseiro, log=lambda m: log(f"F1  {m}"))
N0 = stats['entrada_raw']
N1 = stats['apos_dedup']
_n_dup = stats['duplicados_exatos_removidos']
n_geo_err = stats['erro_geometria']
n_anomalia = stats['anomalia_coleta']
_vd = df.drop_duplicates('CHAVE')['VALIDACAO_ESTAB_DECL'].value_counts()
log(f"T0  sanidade 3 camadas (determinístico+grosseiro=ERRO; percentil=flag): "
    f"ERRO_GEOMETRIA {n_geo_err:,} | FLAG_OUTLIER_GEO {int(df['FLAG_OUTLIER_GEO'].sum()):,}")
log(f"T1  PADRAO_ENDERECO=QUADRA_LOTE: {int((df['PADRAO_ENDERECO']=='QUADRA_LOTE').sum()):,}")
log(f"..  vereditos declarado×contado por endereço: {dict(_vd)}")
log(f"F0.5 anomalia de coleta: {n_anomalia:,} ({n_anomalia/max(N1,1)*100:.1f}% do pós-dedup)")

GATE = rp.gate_universal(df)
GATE_CONF = rp.gate_conf(df)
log(f"GATE aprovados: {int(GATE.sum()):,} / {N1:,} pós-dedup ({N0:,} raw) "
    f"| com faixa ALTA+: {int(GATE_CONF.sum()):,}")

# ── Condomínio horizontal (estrutura física; situação de medição DESCONHECIDA) ──
log("..  detectando condomínios horizontais (estrutura; medição a confirmar no cadastro)...")
df = xu.detectar_condominio_horizontal(df)
log(f"..  registros em condomínio horizontal: {int((df['FLAG_COND_HORIZ'] == 1).sum()):,}")

# ── Agrupamento de coletivas (1 linha por bloco), já gateado ────
log("..  agrupando coletivas por bloco...")
grp = cc.agrupar_enderecos(df)
log(f"..  {len(grp):,} blocos coletivos; {grp['ID_COLETIVA'].nunique():,} endereços")

# ── T5: DBSCAN por densidade de ENDEREÇOS -> polos comerciais ───
log(f"T5  DBSCAN polos comerciais (eps={args.dbscan_eps:.0f}m, min={args.dbscan_min})...")
agg_polo, polos = xu.detectar_polos_comerciais(
    df, eps_m=args.dbscan_eps, min_samples=args.dbscan_min, peso_max=args.dbscan_peso_max)
log(f"T5  {len(polos):,} polos detectados")

# ════════════════════════════════════════════════════════════════
#  COLETIVAS SIMPLES + ECONOMIAS FALTANTES (gap)  — função canônica
# ════════════════════════════════════════════════════════════════
log("..  coletivas simples (2 economias) + economias faltantes (gap)...")
col = df[df['FLAG_COL'] == 1].copy()
qlrec = df[(df['PADRAO_ENDERECO'] == 'QUADRA_LOTE') &
           (df['SANIDADE_GEO'] == 'OK') & (df['FLAG_ANOMALIA_COLETA'] == 0)].copy()
simples, par_faltante, flags_df = xu.construir_simples_e_faltantes(col, ql_records=qlrec)
n_uso_misto = int(flags_df['FLAG_USO_MISTO'].sum())
n_fachada = int(flags_df['FLAG_FACHADA_ATIVA'].sum())
log(f"..  simples (bruto): {len(simples):,} | faltantes: {len(par_faltante):,} | "
    f"uso misto: {n_uso_misto:,} | fachada: {n_fachada:,}")

# ── Condomínios horizontais: TIRAR da classe coletiva → aba própria (individual) ──
# Roteamento SÓ pelo rótulo do classificador (autoritativo horizontal×vertical).
# HORIZ_SUBDIV (fundos/frente/lado/anexo) NÃO entra: 2 economias, ligação compartilhada (alvo).
# casa_seq (FLAG_COND_HORIZ) NÃO roteia (pegaria prédios verticais com registros CASA);
# serve só p/ enriquecer SEQ/GAP nos endereços já rotulados como horizontais.
HORIZ_COND_LABELS = {
    'Vila ou condomínio horizontal (casas)',
    'Condomínio horizontal com múltiplos blocos/vias',
    'Múltiplas moradias no mesmo endereço (casa/lote/sobrado)',
}
ch = df[df['FLAG_COND_HORIZ'] == 1]
if len(ch):
    ch_det = ch.groupby('ID_COLETIVA').agg(
        SEQ_MIN=('COND_HORIZ_SEQ_MIN', 'first'), SEQ_MAX=('COND_HORIZ_SEQ_MAX', 'first'),
        GAP=('COND_HORIZ_GAP', 'first'))
else:
    ch_det = pd.DataFrame(columns=['SEQ_MIN', 'SEQ_MAX', 'GAP'])

# simples (frente/fundos, 2-4 casas no lote) é alvo de ligação compartilhada → permanece
simples_ids = set(simples['ID_COLETIVA'])

# coletivas formais = grp menos simples; depois separa horizontal SÓ por rótulo
flags = flags_df.set_index('ID_COLETIVA')[['FLAG_USO_MISTO', 'FLAG_FACHADA_ATIVA']]
grp = grp[~grp['ID_COLETIVA'].isin(simples_ids)].copy()
grp = grp.merge(flags, left_on='ID_COLETIVA', right_index=True, how='left')
grp['PORTE'] = grp[['N_UNIDADES', 'N_MORADIAS', 'TOTAL_GERAL']].max(axis=1)
# N16: multiplicidade de estabelecimentos DECLARADA pelo recenseador no bloco
grp['MULT_ESTAB_DECL'] = (pd.to_numeric(grp['IND_ESTAB_MAX'], errors='coerce')
                          .map(xu.IND_ESTAB_DESC)
                          if 'IND_ESTAB_MAX' in grp.columns else pd.NA)
is_horiz = grp['TIPO_COLETIVA'].isin(HORIZ_COND_LABELS)
cond_horiz_grp = grp[is_horiz].copy()
grp = grp[~is_horiz].copy()
log(f"..  condomínios horizontais: {len(cond_horiz_grp):,} (rótulo do classificador) -> aba própria; "
    f"coletivas formais (vertical/compartilhada): {len(grp):,}")

# aba Condominios_Horizontais (1 linha por bloco) + detalhe de sequência onde houver
# v4.5: estrutura física NÃO prova situação comercial/hidrométrica — a
# conclusão "provável hidrômetro por casa" extrapolava a evidência. A aba
# afirma só o que sabe (estrutura) e aponta a ação (cruzar com cadastro).
CH_COLS = ['ID_COLETIVA', 'RADAR_ID_ENDERECO', 'RADAR_ID_BLOCO',
           'LOGRADOURO', 'NUMERO', 'LOCALIDADE', 'CEP', 'COD_SETOR',
           'TIPO_COND', 'QTD_CASAS', 'SEQ_MIN', 'SEQ_MAX', 'GAP', 'CONF',
           'N_ECONOMIAS', 'CENTROIDE_LAT', 'CENTROIDE_LON',
           'SITUACAO_MEDICAO', 'ACAO_RECOMENDADA']
if len(cond_horiz_grp):
    cond_horiz = cond_horiz_grp.merge(ch_det, left_on='ID_COLETIVA', right_index=True, how='left')
    cond_horiz['QTD_CASAS'] = cond_horiz[['N_MORADIAS', 'N_UNIDADES', 'TOTAL_GERAL']].max(axis=1)
    cond_horiz['SITUACAO_MEDICAO'] = 'DESCONHECIDA (estrutura nao prova hidrometro)'
    cond_horiz['ACAO_RECOMENDADA'] = 'CRUZAR_COM_CADASTRO_COMERCIAL'
    cond_horiz = cond_horiz.rename(columns={'SETOR': 'COD_SETOR', 'TIPO_COLETIVA': 'TIPO_COND',
                                            'CONFIANCA_COD': 'CONF', 'TOTAL_GERAL': 'N_ECONOMIAS'})
    cond_horiz = cond_horiz.reindex(columns=CH_COLS).sort_values('QTD_CASAS', ascending=False)
else:
    cond_horiz = pd.DataFrame(columns=CH_COLS)

is_alta = grp['CONFIANCA_COD'] == 'ALTA'
is_grande = is_alta & ((grp['N_BLOCOS'] >= 2) | (grp['PORTE'] >= 4))
gr_grandes = grp[is_grande].sort_values(['PORTE', 'TOTAL_GERAL'], ascending=False)
gr_pequenas = grp[grp['CONFIANCA_COD'].isin(['ALTA', 'MEDIA']) & ~is_grande] \
                 .sort_values(['CONFIANCA_COD', 'PORTE'], ascending=[True, False])

# ── ABA "Coletivas": SÓ OBSERVADO (GRANDE/PEQUENA/SIMPLES). v4.5: hipótese
#    NUNCA se mistura a fato — PAR_FALTANTE sai para a aba Hipoteses_Expansao.
#    Schema semântico: QTD_ECONOMIAS (mesma grandeza nas 3 classes) e
#    CONF_TIPOLOGIA (só confiança de tipologia — localização está em F7).
def _harmonizar(dfin, classe, col_tipo, col_qtd, col_conf, fn_det, col_setor='COD_SETOR'):
    o = pd.DataFrame(index=range(len(dfin)))
    o['CLASSE'] = classe
    o['ID_COLETIVA'] = dfin['ID_COLETIVA'].values if 'ID_COLETIVA' in dfin.columns else None
    o['RADAR_ID_ENDERECO'] = (dfin['RADAR_ID_ENDERECO'].values
                              if 'RADAR_ID_ENDERECO' in dfin.columns else None)
    o['RADAR_ID_BLOCO'] = (dfin['RADAR_ID_BLOCO'].values
                           if 'RADAR_ID_BLOCO' in dfin.columns else None)
    o['LOGRADOURO'] = dfin['LOGRADOURO'].values
    o['NUMERO'] = dfin['NUMERO'].values
    o['LOCALIDADE'] = dfin['LOCALIDADE'].values
    o['CEP'] = dfin['CEP'].values
    o['COD_SETOR'] = dfin[col_setor].values
    o['TIPO'] = dfin[col_tipo].values
    o['QTD_ECONOMIAS'] = dfin[col_qtd].values
    o['DETALHE'] = fn_det(dfin).values if len(dfin) else []
    o['FLAG_USO_MISTO'] = dfin['FLAG_USO_MISTO'].values if 'FLAG_USO_MISTO' in dfin.columns else 0
    o['FLAG_FACHADA_ATIVA'] = dfin['FLAG_FACHADA_ATIVA'].values if 'FLAG_FACHADA_ATIVA' in dfin.columns else 0
    o['MULT_ESTAB_DECL'] = (dfin['MULT_ESTAB_DECL'].values
                            if 'MULT_ESTAB_DECL' in dfin.columns else None)
    o['VALIDACAO_ESTAB_DECL'] = (dfin['VALIDACAO_ESTAB_DECL'].values
                                 if 'VALIDACAO_ESTAB_DECL' in dfin.columns else None)
    o['CONF_TIPOLOGIA'] = dfin[col_conf].values
    o['CENTROIDE_LAT'] = dfin['CENTROIDE_LAT'].values
    o['CENTROIDE_LON'] = dfin['CENTROIDE_LON'].values
    return o

_det_gp = lambda d: (d['N_BLOCOS'].fillna(0).astype(int).astype(str) + 'bl/' +
                     d['N_ANDARES'].fillna(0).astype(int).astype(str) + 'and/' +
                     d['N_UNIDADES'].fillna(0).astype(int).astype(str) + 'un/' +
                     d['N_MORADIAS'].fillna(0).astype(int).astype(str) + 'casas')
_det_si = lambda d: d['MARCADORES'].astype(str)

_frames_col = [
    _harmonizar(gr_grandes, 'GRANDE', 'TIPO_COLETIVA', 'PORTE', 'CONFIANCA_COD', _det_gp, 'SETOR'),
    _harmonizar(gr_pequenas, 'PEQUENA', 'TIPO_COLETIVA', 'PORTE', 'CONFIANCA_COD', _det_gp, 'SETOR'),
    _harmonizar(simples, 'SIMPLES', 'TIPO_SIMPLES', 'QTD_ECONOMIAS', 'CONFIANCA_COD', _det_si, 'COD_SETOR'),
]
# concat só de frames não-vazios (evita FutureWarning de colunas all-NA)
_nv = [f for f in _frames_col if len(f)]
coletivas = pd.concat(_nv if _nv else _frames_col[:1], ignore_index=True)
_co = {'GRANDE': 0, 'PEQUENA': 1, 'SIMPLES': 2}
_cf = {'MUITO_ALTA': 0, 'ALTA': 0, 'MEDIA': 1, 'BAIXA': 2}
coletivas['_o'] = coletivas['CLASSE'].map(_co)
coletivas['_f'] = coletivas['CONF_TIPOLOGIA'].map(lambda x: _cf.get(str(x), 3))
coletivas = coletivas.sort_values(['_o', '_f', 'QTD_ECONOMIAS'],
                                  ascending=[True, True, False]).drop(columns=['_o', '_f'])
n_grande = int((coletivas['CLASSE'] == 'GRANDE').sum())
n_pequena = int((coletivas['CLASSE'] == 'PEQUENA').sum())
log(f"..  COLETIVAS (observado): {len(coletivas):,} "
    f"[G {n_grande} / P {n_pequena} / S {len(simples)}]")

# ── ABA "Hipoteses_Expansao" (v4.5): a inferência de gap NUNCA se fantasia
#    de fato observado — sai com semântica explícita de hipótese.
HIP_COLS = ['RADAR_ID_ENDERECO', 'ID_COLETIVA', 'LOGRADOURO', 'NUMERO',
            'LOCALIDADE', 'CEP', 'COD_SETOR', 'TIPO_HIPOTESE', 'EVIDENCIA',
            'UNIDADE_INFERIDA', 'N_AUSENTES', 'CONF_INFERENCIA',
            'REGRA_ORIGEM', 'REQUER_CONFIRMACAO', 'RADAR_CONF_COORD',
            'CENTROIDE_LAT', 'CENTROIDE_LON']
_REGRA = {'POSICIONAL': 'gap_posicional', 'SEQUENCIA': 'gap_linear',
          'GRADE_APTO': 'gap_grade', 'QUADRA_LOTE': 'gap_quadra_lote'}
if len(par_faltante):
    hipoteses = par_faltante.rename(columns={
        'GAP_TIPO': 'TIPO_HIPOTESE', 'PRESENTES': 'EVIDENCIA',
        'AUSENTE_INFERIDO': 'UNIDADE_INFERIDA', 'CONF_GAP': 'CONF_INFERENCIA'}).copy()
    hipoteses['REGRA_ORIGEM'] = hipoteses['TIPO_HIPOTESE'].map(_REGRA)
    hipoteses['REQUER_CONFIRMACAO'] = 'SIM'
    hipoteses = hipoteses.reindex(columns=HIP_COLS)
else:
    hipoteses = pd.DataFrame(columns=HIP_COLS)
log(f"..  HIPOTESES_EXPANSAO (inferidas, requerem confirmação): {len(hipoteses):,}")

# Não-residencial (potencial comercial), gateado
# v4.3: FLAG_USO_MISTO por ENDEREÇO (set-based) — res + não-res na mesma CHAVE
_esp = pd.to_numeric(df['COD_ESPECIE'], errors='coerce')
_mix = (df.assign(_r=_esp.isin([1, 2]).astype(int),
                  _n=_esp.isin([3, 4, 5, 6, 8]).astype(int))
          .groupby('CHAVE')[['_r', '_n']].transform('max'))
df['FLAG_USO_MISTO'] = ((_mix['_r'] == 1) & (_mix['_n'] == 1)).astype(int)

# v4.5: F7 é gate — Nao_Residencial exige faixa de localização ALTA+
nres = df[GATE_CONF & ~df['SETOR_ATIVIDADE'].isin(['RESIDENCIAL', 'CONSTRUCAO'])].copy()
NRES_COLS = ['COD_UNICO_ENDERECO', 'RADAR_ID_ENDERECO', 'DSC_LOCALIDADE',
             'LOGRADOURO', 'NUMERO',
             'COMPLEMENTO_NORM', 'CEP', 'COD_SETOR', 'LATITUDE', 'LONGITUDE',
             'NV_GEO_COORD', 'QUALIDADE_COORD', 'SETOR_ATIVIDADE', 'ATIVIDADE_DETALHE',
             'DSC_ESTABELECIMENTO', 'DSC_ESPECIE', 'FLAG_USO_MISTO',
             'DSC_MULT_ESTAB', 'QTD_ESTAB_CONTADA', 'VALIDACAO_ESTAB_DECL',
             'RADAR_CONF_LOCALIZACAO', 'RADAR_CONF_FAIXA']
nres = nres[NRES_COLS].sort_values(
    ['RADAR_CONF_LOCALIZACAO', 'SETOR_ATIVIDADE'], ascending=[False, True])

# Fraude: detectada e EXCLUÍDA do gate (R2 preservado). Aba de auditoria removida do output.

# Polos
POLO_COLS = ['POLO_ID', 'QTD_END', 'QTD_ESTAB', 'SETOR_DOMINANTE', 'SETORES',
             'CENTROIDE_LAT', 'CENTROIDE_LON', 'CLASSE_POLO', 'POLIGONO_WKT']
polos_out = polos[POLO_COLS].copy() if len(polos) else pd.DataFrame(columns=POLO_COLS)

# ── Padronizar TODAS as colunas de coordenadas (lat e lon) em 6 casas decimais (~0,11 m) ──
_COORD_COLS = ['CENTROIDE_LAT', 'CENTROIDE_LON', 'LATITUDE', 'LONGITUDE']
for _dfo in (coletivas, cond_horiz, nres, polos_out, hipoteses):
    for _c in _COORD_COLS:
        if _c in _dfo.columns:
            _dfo[_c] = pd.to_numeric(_dfo[_c], errors='coerce').round(6)

# ── RADAR_ID_* como TEXTO no Excel: 19 dígitos > 2^53 estouram o float64 da
#    célula (perda de precisão silenciosa). Em banco/Parquet, usar BIGINT.
_ID_COLS = ('RADAR_ID_ENDERECO', 'RADAR_ID_BLOCO')
for _dfo in (coletivas, cond_horiz, nres, hipoteses):
    for _c in _ID_COLS:
        if _c in _dfo.columns:
            # NUNCA passar por float64 (to_numeric coagiria e perderia dígitos)
            _dfo[_c] = pd.Series(pd.array(_dfo[_c], dtype='Int64'),
                                 index=_dfo.index).astype(str).replace({'<NA>': ''})

# ── distribuição p/ o print final ───────────────────────────────
faixa_dist = (df[GATE]['RADAR_CONF_FAIXA'].value_counts()
              .reindex(['MUITO_ALTA', 'ALTA', 'MEDIA', 'BAIXA']).fillna(0).astype(int))

# ════════════════════════════════════════════════════════════════
#  F7b — GATE EXECUTÁVEL (M2): roda ANTES de qualquer export.
#  Violação de R1/R2/R3/H/R5 -> RadarQualityError (imune a python -O).
# ════════════════════════════════════════════════════════════════
log("F7b gate_pos_condicoes (R1/R2/R3/H/R5, RadarQualityError)...")
intro, dic = xu.construir_dicionario()
_ABAS = ('Dicionario', 'Coletivas', 'Hipoteses_Expansao',
         'Condominios_Horizontais', 'Nao_Residencial', 'Polos_Comerciais')
xu.gate_pos_condicoes(coletivas=coletivas, cond_horiz=cond_horiz, nres=nres,
                      polos=polos_out, hipoteses=hipoteses, df_base=df,
                      dicionario=dic, nomes_abas=_ABAS)
log("F7b gate OK — export liberado")

# ════════════════════════════════════════════════════════════════
#  EXCEL
# ════════════════════════════════════════════════════════════════
log("XLS gerando Excel de potenciais...")
wb = Workbook()
ws = wb.active
ws.title = 'Dicionario'
ws.cell(1, 1, 'DICIONÁRIO DE DADOS — Radar Coletivo (Potenciais)').font = cc._TF
r = 3
for ln in intro:
    cell = ws.cell(r, 1, '• ' + ln)
    cell.font = cc._SF
    cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 30
    r += 1
r += 1

# cabeçalho da tabela de colunas
hdr = r
_HF = PatternFill('solid', fgColor='D9E1F2')
for j, h in enumerate(['ABA', 'COLUNA', 'TIPO', 'DESCRIÇÃO', 'VALORES / DOMÍNIO'], 1):
    c = ws.cell(hdr, j, h)
    c.font = cc._BF
    c.fill = _HF
r += 1

# uma linha por coluna documentada, com listra por aba
_ABA_FILL = PatternFill('solid', fgColor='EEF3FB')
abas_ordem = list(dict.fromkeys(dic['ABA']))
for _, row in dic.iterrows():
    listra = _ABA_FILL if abas_ordem.index(row['ABA']) % 2 == 1 else None
    vals = [row['ABA'], row['COLUNA'], row['TIPO'], row['DESCRICAO'], row['VALORES']]
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v)
        if j in (1, 2):
            c.font = cc._BF
        if j in (4, 5):
            c.alignment = Alignment(wrap_text=True, vertical='top')
        else:
            c.alignment = Alignment(vertical='top')
        if listra is not None:
            c.fill = listra
    r += 1

ws.freeze_panes = f'A{hdr + 1}'
for col, w in [('A', 24), ('B', 26), ('C', 14), ('D', 62), ('E', 72)]:
    ws.column_dimensions[col].width = w

# abas de potenciais (formatação do runner)
cc._write_sheet(wb, 'Coletivas',
                f'COLETIVAS OBSERVADAS — {len(coletivas):,} (CLASSE: GRANDE/PEQUENA/SIMPLES)',
                '100% observado em campo — hipóteses inferidas estão na aba Hipoteses_Expansao. '
                'QTD_ECONOMIAS é grandeza única; CONF_TIPOLOGIA = confiança da classificação. Gate R1-R3.',
                coletivas, conf_col='CONF_TIPOLOGIA')
cc._write_sheet(wb, 'Hipoteses_Expansao',
                f'HIPÓTESES DE EXPANSÃO — {len(hipoteses):,} (INFERIDAS, requerem confirmação)',
                'Unidades AUSENTES inferidas por gap analysis: são HIPÓTESE, não fato — a lacuna pode '
                'ser praça, lote unido ou numeração descontinuada. RADAR_CONF_COORD=0 sempre; '
                'REQUER_CONFIRMACAO=SIM. Nunca somar com as abas observadas.',
                hipoteses, conf_col='CONF_INFERENCIA')
cc._write_sheet(wb, 'Condominios_Horizontais',
                f'CONDOMÍNIOS HORIZONTAIS — {len(cond_horiz):,} endereços (estrutura horizontal)',
                'Casas/lotes (3+) em condomínio horizontal. A estrutura física NÃO prova a situação '
                'hidrométrica: SITUACAO_MEDICAO=DESCONHECIDA; ação = cruzar com o cadastro comercial '
                '(6 casas/1 ligação = oportunidade; 6/6 = individualizado).',
                cond_horiz, conf_col='CONF')
cc._write_sheet(wb, 'Nao_Residencial',
                f'NÃO-RESIDENCIAL — {len(nres):,} (potencial comercial, faixa ALTA+)',
                'Espécies != residencial/construção, com RADAR_CONF_FAIXA ALTA/MUITO_ALTA (F7 é gate). '
                'Ordenado por RADAR_CONF_LOCALIZACAO.',
                nres, conf_col='RADAR_CONF_FAIXA', setor_col='SETOR_ATIVIDADE')
cc._write_sheet(wb, 'Polos_Comerciais',
                f'POLOS COMERCIAIS — {len(polos_out):,} (DBSCAN por densidade de ENDEREÇOS)',
                'min_samples = endereços (sem peso — v4.5); intensidade comercial (QTD_ESTAB) '
                'caracteriza o cluster. CLASSE_POLO por porte; POLIGONO_WKT (EPSG:4326) p/ QGIS.',
                polos_out)

# Exibição consistente: 6 casas decimais nas colunas de coordenadas (todas as abas)
_COORD_HEADERS = {'CENTROIDE_LAT', 'CENTROIDE_LON', 'LATITUDE', 'LONGITUDE'}
for _ws in wb.worksheets:
    _hdr_row, _cols = None, []
    for _row in _ws.iter_rows(min_row=1, max_row=6):
        _names = [c.value for c in _row]
        if any(n in _COORD_HEADERS for n in _names):
            _hdr_row = _row[0].row
            _cols = [c.column for c in _row if c.value in _COORD_HEADERS]
            break
    if _hdr_row:
        for _ci in _cols:
            for _ri in range(_hdr_row + 1, _ws.max_row + 1):
                _ws.cell(_ri, _ci).number_format = '0.000000'

wb.save(args.output)
log(f"XLS salvo: {args.output}")
log(f"    abas: {wb.sheetnames}")

# ════════════════════════════════════════════════════════════════
#  MANIFEST DE LINHAGEM (v4.5): R5 remove a MARCA da fonte da
#  apresentação comercial; a linhagem TÉCNICA não pode ser destruída.
#  Sidecar JSON com proveniência, parâmetros e FUNIL de exclusões —
#  fora do xlsx comercial, fora do escopo do scan R5 (por design).
# ════════════════════════════════════════════════════════════════
import json as _json
import hashlib as _hashlib
from datetime import datetime as _dt


def _sha256_arquivo(caminho, buf=1 << 20):
    h = _hashlib.sha256()
    with open(caminho, 'rb') as f:
        while True:
            b = f.read(buf)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


# ── Fonte: declarada > inferida por nome > UNKNOWN. NUNCA inventada
# (v4.5.1: antes era constante hard-coded — provenance falsa).
_src_name = os.path.basename(args.input)
if args.source_dataset:
    _src_ds, _src_vin = args.source_dataset, (args.source_vintage or 'UNKNOWN')
    _src_mode = 'declarado_por_argumento'
elif re.search(r'(?i)cnefe.{0,3}2022', _src_name):
    _src_ds, _src_vin = 'IBGE CNEFE Censo 2022', '2022'
    _src_mode = 'inferido_por_nome_de_arquivo'
else:
    _src_ds, _src_vin, _src_mode = 'UNKNOWN', 'UNKNOWN', 'nao_identificado'

# ── Funil por CAUSA PRIORITÁRIA (NUM=0 > geometria > anomalia): cada
# registro excluído conta UMA vez; o funil fecha por construção. As
# sobreposições brutas saem à parte (um registro pode ter as 3 causas).
_c_num = df['NUMERO'] <= 0
_c_geo = df['SANIDADE_GEO'] != 'OK'
# PATCH LOCAL (execucao Corsan/RS) — a causa de exclusao por anomalia era
# contada pela flag BINARIA enquanto o gate ja aplicava a ESCADA de sinais
# primarios (v4.11.0). O funil deixava de fechar: a linha 'consistencia' saia
# aritmeticamente falsa no manifest (ex.: "505 - 9 - 1 - 38 = 495"), somando
# como excluido quem o gate deixou passar. W6 exige que o funil feche por
# construcao — a causa passa a ser a que o motor EXECUTA.
_lim_ano = rp.ANOMALIA_EXCLUI_A_PARTIR_DE
_sin_prim = (df['N_SINAIS_PRIMARIOS'] if 'N_SINAIS_PRIMARIOS' in df.columns
             else pd.Series(0, index=df.index)).fillna(0).astype(int)
_c_ano = _sin_prim >= _lim_ano
_c_ano_flag = df['FLAG_ANOMALIA_COLETA'] == 1
_ex_num = int(_c_num.sum())
_ex_geo = int((~_c_num & _c_geo).sum())
_ex_ano = int((~_c_num & ~_c_geo & _c_ano).sum())

import platform
from importlib import metadata as _ilmd


def _validar_runtime_vs_lock():
    """Compara o runtime corrente com requirements.lock.txt (3ª revisão):
    a execução declara se rodou no ambiente EXATAMENTE validado. Divergência
    NÃO aborta — vira WARNING_RUNTIME_DIVERGENTE gravado na linhagem."""
    lock_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                             'requirements.lock.txt')
    if not os.path.exists(lock_path):
        return {'status': 'LOCK_AUSENTE', 'lock_match': None, 'lock_sha256': None}
    with open(lock_path, 'rb') as f:
        lock_sha = _hashlib.sha256(f.read()).hexdigest()
    pins = {}
    with open(lock_path, encoding='utf-8') as f:
        for ln in f:
            ln = ln.strip()
            if ln and not ln.startswith('#') and '==' in ln:
                p, v = ln.split('==', 1)
                pins[p.strip().lower()] = v.strip()
    alvo = ('pandas', 'numpy', 'scikit-learn', 'openpyxl')
    atuais = {p: _ilmd.version(p) for p in alvo}
    diverg = {p: {'lock': pins.get(p), 'runtime': atuais[p]}
              for p in alvo if pins.get(p) != atuais[p]}
    return {
        'lock_sha256': lock_sha,
        'lock_match': not diverg,
        'divergencias': diverg,
        'status': 'VALIDATED' if not diverg else 'WARNING_RUNTIME_DIVERGENTE',
    }


manifest = {
    # 4.5.3 — R8 (integridade de identificadores) + FK hipótese→pai. A versão
    # é o que a forense usa para separar entrega afetada de entrega sã.
    'pipeline': 'radar_coletivo', 'pipeline_version': xu.VERSAO,
    'run_executed_at': _dt.now().astimezone().isoformat(timespec='seconds'),
    'source': {
        'dataset': _src_ds, 'vintage': _src_vin,
        'file': _src_name, 'sha256': _sha256_arquivo(args.input),
        'metadata_mode': _src_mode,
        'verified': False,   # True só com sidecar assinado do downloader (roadmap)
    },
    'runtime': {
        'python': platform.python_version(),
        **{p: _ilmd.version(p) for p in
           ('pandas', 'numpy', 'scikit-learn', 'openpyxl')},
    },
    'runtime_validation': _validar_runtime_vs_lock(),
    'parameters': {k: v for k, v in vars(args).items()},
    'funil': {
        'entrada_raw': int(N0),
        'duplicados_exatos_removidos': int(_n_dup),
        'apos_dedup': int(N1),
        'exclusoes_causa_prioritaria': {
            'numero_invalido': _ex_num,
            'geometria_invalida': _ex_geo,
            'anomalia_coleta': _ex_ano,
        },
        'gate_universal_aprovado': int(GATE.sum()),
        'gate_f7_alta_mais': int(GATE_CONF.sum()),
        'consistencia': f'{N1} - {_ex_num} - {_ex_geo} - {_ex_ano} = {int(GATE.sum())}',
        'sobreposicao_causas_bruta': {
            'numero_e_geometria': int((_c_num & _c_geo).sum()),
            'numero_e_anomalia': int((_c_num & _c_ano).sum()),
            'geometria_e_anomalia': int((_c_geo & _c_ano).sum()),
            'todas': int((_c_num & _c_geo & _c_ano).sum()),
        },
        # A escada RETEM o registro de 1-2 sinais primarios. Sem esta linha,
        # quem le o funil nao distingue "nao houve anomalia" de "houve e foi
        # deixada passar" — que e exatamente a decisao que a v4.11.0 tomou.
        # Enderecos cujas CHAVEs operacionais distintas cairam no MESMO id
        # por normalizacao declarada (numeral por extenso, titulo omitido).
        # E fusao pretendida — e precisa estar visivel para quem consome.
        'n14_fusoes_por_normalizacao': int(
            getattr(xu, 'N14_FUSOES_POR_NORMALIZACAO', 0)),
        'anomalia_escada': {
            'limiar_exclusao_sinais_primarios': int(_lim_ano),
            'marcada_pela_flag_binaria': int(_c_ano_flag.sum()),
            'excluida_pela_escada': int(_c_ano.sum()),
            'retida_com_observacao_1_sinal': int((_sin_prim == 1).sum()),
            'retida_com_quarentena_2_sinais': int((_sin_prim == 2).sum()),
        },
        'flag_outlier_geo_informativo': int(df.get('FLAG_OUTLIER_GEO', pd.Series(0, index=df.index)).sum()),
        'numero_zero_com_chave_ql_elegivel_hipotese': int(
            (_c_num & (df['PADRAO_ENDERECO'] == 'QUADRA_LOTE') & ~_c_geo & ~_c_ano).sum()),
    },
    'saidas': {
        'coletivas_observadas': int(len(coletivas)),
        'hipoteses_expansao': int(len(hipoteses)),
        'condominios_horizontais': int(len(cond_horiz)),
        'nao_residencial': int(len(nres)),
        'polos_comerciais': int(len(polos_out)),
    },
    'faixas_confiabilidade_gate': {k: int(v) for k, v in faixa_dist.items()},
    'vereditos_declarado_x_contado': {str(k): int(v) for k, v in _vd.items()},
    'avisos': [
        'RADAR_ID e deterministico da REPRESENTACAO canonica, nao identidade '
        'fisica permanente — mudanca de nome de logradouro/localidade muda o id.',
        'RADAR_CONF_* e score heuristico de completude de evidencia; '
        'calibracao contra verdade externa (M6) pendente.',
        'Hipoteses_Expansao sao inferencias — requerem confirmacao de campo.',
        'Determinismo e SEMANTICO (valores identicos entre execucoes); o '
        'SHA-256 do .xlsx varia por metadados internos do formato.',
        'Raio grosseiro de geometria (300 km, ajustavel) e divida aberta — '
        'substituir por malha municipal ST_Covers + buffer.',
    ],
}
_manifest_path = args.output + '.manifest.json'
with open(_manifest_path, 'w', encoding='utf-8') as f:
    _json.dump(manifest, f, ensure_ascii=False, indent=2)
log(f"MANIFEST de linhagem: {_manifest_path}")

print("\nRESUMO FINAL")
print(f"  COLETIVAS observadas : {len(coletivas):,}  [G {n_grande} / P {n_pequena} / S {len(simples)}]")
print(f"  Simples por tipo     : {dict(simples['TIPO_SIMPLES'].value_counts())}")
print(f"  HIPÓTESES (inferidas): {len(hipoteses):,}  {dict(hipoteses['TIPO_HIPOTESE'].value_counts()) if len(hipoteses) else {}}")
print(f"  Cond. horizontais    : {len(cond_horiz):,}  (situação de medição: DESCONHECIDA — cruzar c/ cadastro)")
print(f"  Nao-residencial      : {len(nres):,} (faixa ALTA+)")
print(f"  Polos comerciais     : {len(polos_out):,}  -> {dict(polos_out['CLASSE_POLO'].value_counts()) if len(polos_out) else {}}")
print(f"  Anomalia de coleta excluída (gate, R2): {n_anomalia:,}")
print(f"  RADAR_CONF (gate)    : {dict(faixa_dist)}")
