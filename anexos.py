# -*- coding: utf-8 -*-
"""Transforma o que o usuário anexa em algo que o modelo consegue ler.

DUAS ROTAS, E ELAS SÃO MESMO DIFERENTES

  IMAGEM  → vai INTEIRA para o modelo, em base64. O `qwen3vl` é um modelo de
            visão: ele enxerga a foto, não uma descrição dela. O servidor subiu
            com `--limit-mm-per-prompt.image 30`, então cabem até 30 por
            mensagem.

  ARQUIVO → vira TEXTO antes de entrar. PDF, Word, planilha e ZIP não são
            imagens; mandá-los em base64 gastaria contexto para nada, porque o
            modelo não decodifica formato binário.

Um PDF de fotos escaneadas cai num vão entre as duas: o texto extraído vem
vazio e ninguém avisa. Por isso a extração devolve `paginas_sem_texto` — quem
consome decide se manda as páginas como imagem.

Limites existem para proteger a janela de contexto: um relatório de 200 páginas
entra inteiro e empurra a conversa para fora, e o modelo passa a responder sobre
o começo do arquivo em vez da sua pergunta.
"""
from __future__ import annotations

import base64
import io
import json
import pathlib
import re
import zipfile

# ─── limites ─────────────────────────────────────────────────────────────────
MAX_BYTES = 25 * 1024 * 1024        # por arquivo
MAX_TEXTO = 40_000                  # caracteres por documento
MAX_ITENS_ZIP = 40                  # arquivos listados/lidos de um ZIP
MAX_LADO = 1600                     # px: imagem maior é reduzida antes de ir

IMAGENS = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp"}
TEXTOS = {".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm", ".log",
          ".sql", ".py", ".js", ".ts", ".yaml", ".yml", ".ini", ".env"}


def tipo_de(nome: str) -> str:
    ext = pathlib.Path(nome).suffix.lower()
    if ext in IMAGENS:
        return "imagem"
    if ext == ".pdf":
        return "pdf"
    if ext in (".docx", ".doc"):
        return "word"
    if ext in (".xlsx", ".xlsm", ".xls"):
        return "planilha"
    if ext == ".zip":
        return "zip"
    if ext in TEXTOS:
        return "texto"
    return "desconhecido"


def _cortar(txt: str, limite: int = MAX_TEXTO) -> tuple[str, bool]:
    t = re.sub(r"\n{3,}", "\n\n", (txt or "").strip())
    return (t[:limite], len(t) > limite)


# ─── imagem ──────────────────────────────────────────────────────────────────
def preparar_imagem(dados: bytes, nome: str) -> dict:
    """Reduz se for grande demais e devolve em base64, pronto para o modelo.

    Reduzir não é economia de disco — é de CONTEXTO. Uma foto de 4000 px vira
    milhares de tokens visuais e não acrescenta nada além do que 1600 px já
    mostram para ler fachada ou documento.
    """
    from PIL import Image

    try:
        img = Image.open(io.BytesIO(dados))
        img.load()
    except Exception as e:
        return {"erro": f"não consegui abrir a imagem: {type(e).__name__}"}

    largura, altura = img.size
    reduzida = False
    if max(largura, altura) > MAX_LADO:
        escala = MAX_LADO / max(largura, altura)
        img = img.resize((int(largura * escala), int(altura * escala)),
                         Image.LANCZOS)
        reduzida = True
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")

    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    return {
        "tipo": "imagem", "nome": nome,
        "dimensao": f"{largura}x{altura}",
        "reduzida": reduzida,
        "b64": base64.b64encode(buf.getvalue()).decode(),
    }


# ─── documentos ──────────────────────────────────────────────────────────────
def _pdf(dados: bytes) -> dict:
    from pypdf import PdfReader

    leitor = PdfReader(io.BytesIO(dados))
    partes, vazias = [], 0
    for i, pag in enumerate(leitor.pages, 1):
        try:
            t = (pag.extract_text() or "").strip()
        except Exception:
            t = ""
        if t:
            partes.append(f"--- página {i} ---\n{t}")
        else:
            vazias += 1
    texto, cortado = _cortar("\n\n".join(partes))
    return {"paginas": len(leitor.pages), "paginas_sem_texto": vazias,
            "texto": texto, "cortado": cortado}


def _word(dados: bytes) -> dict:
    import docx

    d = docx.Document(io.BytesIO(dados))
    linhas = [p.text for p in d.paragraphs if p.text.strip()]
    # tabelas do Word somem quando se lê só parágrafos, e é comum o dado que
    # interessa estar exatamente nelas
    for t in d.tables:
        for linha in t.rows:
            celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
            if celulas:
                linhas.append(" | ".join(celulas))
    texto, cortado = _cortar("\n".join(linhas))
    return {"paragrafos": len(d.paragraphs), "tabelas": len(d.tables),
            "texto": texto, "cortado": cortado}


def _planilha(dados: bytes) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(dados), read_only=True,
                                data_only=True)
    partes = []
    for aba in wb.worksheets:
        partes.append(f"--- aba {aba.title} ---")
        for i, linha in enumerate(aba.iter_rows(values_only=True)):
            if i >= 200:            # 200 linhas por aba dão a forma do dado
                partes.append("… (restante omitido)")
                break
            vals = [str(v) for v in linha if v is not None]
            if vals:
                partes.append(" | ".join(vals))
    texto, cortado = _cortar("\n".join(partes))
    return {"abas": [a.title for a in wb.worksheets], "texto": texto,
            "cortado": cortado}


def _zip(dados: bytes) -> dict:
    """Lista o conteúdo e lê os arquivos de texto que couberem.

    Não extrai imagem de dentro do ZIP: para vê-las o usuário anexa a imagem, e
    misturar as duas rotas aqui esconderia de quem chama que houve foto dentro.
    """
    z = zipfile.ZipFile(io.BytesIO(dados))
    itens = [i for i in z.infolist() if not i.is_dir()]
    lista = [{"nome": i.filename, "bytes": i.file_size}
             for i in itens[:MAX_ITENS_ZIP]]
    partes = []
    for i in itens[:MAX_ITENS_ZIP]:
        if tipo_de(i.filename) not in ("texto", "pdf") or i.file_size > 2_000_000:
            continue
        try:
            bruto = z.read(i)
        except Exception:
            continue
        if tipo_de(i.filename) == "pdf":
            t = _pdf(bruto).get("texto", "")
        else:
            t = bruto.decode("utf-8", "replace")
        if t.strip():
            partes.append(f"=== {i.filename} ===\n{t[:6000]}")
    texto, cortado = _cortar("\n\n".join(partes))
    return {"arquivos": len(itens), "listagem": lista, "texto": texto,
            "cortado": cortado, "omitidos": max(0, len(itens) - MAX_ITENS_ZIP)}


LEITORES = {"pdf": _pdf, "word": _word, "planilha": _planilha, "zip": _zip}


def ler(dados: bytes, nome: str) -> dict:
    """Extrai o que o modelo vai ler. Devolve sempre `tipo` e `nome`."""
    if len(dados) > MAX_BYTES:
        return {"tipo": "erro", "nome": nome,
                "erro": f"arquivo de {len(dados)//1024//1024} MB; o teto é "
                        f"{MAX_BYTES//1024//1024} MB"}
    t = tipo_de(nome)
    if t == "imagem":
        return preparar_imagem(dados, nome)
    if t == "texto":
        texto, cortado = _cortar(dados.decode("utf-8", "replace"))
        return {"tipo": "texto", "nome": nome, "texto": texto,
                "cortado": cortado}
    leitor = LEITORES.get(t)
    if not leitor:
        return {"tipo": "erro", "nome": nome,
                "erro": f"não sei ler '{pathlib.Path(nome).suffix}'"}
    try:
        d = leitor(dados)
    except Exception as e:
        return {"tipo": "erro", "nome": nome,
                "erro": f"{type(e).__name__}: {str(e)[:200]}"}
    d.update({"tipo": t, "nome": nome})
    return d


def para_mensagem(anexos: list, pergunta: str) -> dict:
    """Monta a mensagem do usuário com texto e imagens juntos.

    Imagens vão como `image_url`, no formato multimodal da API. Documentos
    entram como texto ANTES da pergunta, para o modelo ler o material e só
    então o pedido — a ordem inversa faz ele responder antes de ler.
    """
    blocos, imagens, avisos = [], [], []
    for a in anexos:
        if a.get("tipo") == "imagem" and a.get("b64"):
            imagens.append({"type": "image_url", "image_url": {
                "url": "data:image/jpeg;base64," + a["b64"]}})
            avisos.append(f"[imagem anexada: {a['nome']} · {a.get('dimensao')}"
                          + (" · reduzida" if a.get("reduzida") else "") + "]")
        elif a.get("tipo") == "erro":
            avisos.append(f"[{a['nome']}: {a['erro']}]")
        elif a.get("texto"):
            cab = f"=== {a['nome']} ({a['tipo']}) ==="
            extra = []
            if a.get("paginas"):
                extra.append(f"{a['paginas']} páginas")
            if a.get("paginas_sem_texto"):
                # o caso do PDF escaneado: dizer em vez de entregar vazio
                extra.append(f"{a['paginas_sem_texto']} SEM texto extraível "
                             "(possível PDF de imagem)")
            if a.get("arquivos"):
                extra.append(f"{a['arquivos']} arquivos")
            if a.get("cortado"):
                extra.append("TRUNCADO")
            if extra:
                cab += "  [" + " · ".join(extra) + "]"
            blocos.append(cab + "\n" + a["texto"])
        else:
            avisos.append(f"[{a.get('nome')}: sem texto extraível]")

    partes = blocos + ([""] if blocos else []) + avisos + [pergunta]
    texto = "\n\n".join(p for p in partes if p)
    if not imagens:
        return {"role": "user", "content": texto}
    return {"role": "user",
            "content": [{"type": "text", "text": texto}] + imagens}
