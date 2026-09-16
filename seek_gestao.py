# -*- coding: utf-8 -*-
"""seek_gestao.py — a gestao das aprovacoes da SEEK (14/09/2026).

Pedido do dono do produto: uma pagina para administrador para cima ver quem decide
o que, em quanto tempo, ligacao por ligacao, e exportar — com cada exportacao
registrada. O administrador ve a propria empresa; o suporte (root) ve todas e
escolhe.

    GET  /gestao                                  a pagina
    GET  /api/seek/gestao/resumo                  usuarios: status, taxa, tempo, ultima atividade
    GET  /api/seek/gestao/ligacoes                as decisoes de um usuario (ou de todos), paginadas
    GET  /api/seek/gestao/ligacao/{ligacao}       o historico de uma ligacao: decisoes, comentarios, aberturas
    POST /api/seek/gestao/exportar                XLSX ou CSV, registrado em `seek_exportacao`
    GET  /api/seek/gestao/exportacoes             as exportacoes feitas

TUDO PELA CONEXAO DO USUARIO: a RLS isola a empresa; o nivel e conferido aqui e,
no registro das exportacoes, tambem na politica.

OS CRUZAMENTOS SAO EM PYTHON, por conjunto (ver `nao-usar-sql-caro-tem-py`): as
decisoes do periodo vem por (id_empresa, em); o cadastro e o veredito da IA, pela
chave, so das ligacoes que apareceram.

VIGENTE e a decisao mais recente da ligacao. Numa ligacao decidida duas vezes, a
primeira conta no total de decisoes de quem a tomou, mas nao nas vigentes.

ONDE QUEBRA SOB CARGA: o resumo traz para a API toda decisao do periodo — ~500 mil
linhas por pedido e o limite de um processo com 1 GB. Acima disso o resumo vira
agregacao pronta (tabela diaria por empresa/pessoa/acao, alimentada fora do
Postgres que atende usuario) e a pagina le a agregacao. A exportacao escreve o
arquivo no /tmp do conteiner (tmpfs, conta na memoria): por isso o teto de
`MAX_LINHAS`; acima dele, gerar em fila de trabalho e entregar pelo Storage.
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
import statistics
import tempfile
import time
import uuid
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel

import auth as _auth

router = APIRouter()

FRONT = Path(__file__).resolve().parent / "frontend"
ACOES = ("aprovar", "campo", "revisar", "rejeitar")
ROTULO_ACAO = {"aprovar": "aprovada", "campo": "mandada a campo", "revisar": "em revisão", "rejeitar": "rejeitada"}
ROTULO_IA = {"aprovado": "aprovado", "revisao_humana": "revisão humana", "reprovado": "reprovado"}
#: O Brasil nao tem horario de verao desde 2019: o dia do filtro e o de Brasilia.
FUSO = timezone(timedelta(hours=-3))
#: Teto de linhas por exportacao (ver o docstring: o arquivo nasce no tmpfs).
MAX_LINHAS = 100_000
#: Teto de linhas varridas por pagina da lista de ligacoes com filtro feito em Python.
MAX_VARRIDAS = 50_000

COLUNAS = (
    ("ligacao", "Ligação"),
    ("empresa", "Empresa"),
    ("cidade", "Cidade"),
    ("bairro", "Bairro"),
    ("endereco", "Endereço"),
    ("titular", "Titular"),
    ("qualificacao", "Qualificação no cadastro"),
    ("veredito_ia", "Veredito da IA"),
    ("motivo_ia", "Motivo da IA"),
    ("julgada_em", "Julgada pela IA em"),
    ("decisao", "Decisão humana"),
    ("comentario", "Comentário"),
    ("quem", "Quem decidiu"),
    ("quando", "Quando"),
    ("aberta_em", "Aberta em"),
    ("tempo_avaliacao", "Tempo de avaliação"),
    ("tempo_avaliacao_s", "Tempo de avaliação (s)"),
    ("em_lote", "Em lote"),
    ("vigente", "É a decisão vigente"),
    ("decisao_vigente", "Decisão vigente hoje"),
)
NOMES_COLUNAS = dict(COLUNAS)
COLUNAS_PADRAO = ("ligacao", "cidade", "endereco", "veredito_ia", "motivo_ia", "decisao", "comentario",
                  "quem", "quando", "tempo_avaliacao")


# ─────────────────────────────────────────────────────────── comum ──
def _quem() -> _auth.Usuario:
    u = _auth.USUARIO_DA_REQUISICAO.get()
    if u is None:
        raise HTTPException(401, "sessão ausente")
    if not u.pode("admin"):
        raise HTTPException(403, "a gestão das aprovações exige nível administrador ou acima")
    return u


def _dia(txt, padrao):
    if not txt:
        return padrao
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(txt)):
        raise HTTPException(400, "data inválida: use AAAA-MM-DD")
    try:
        return date.fromisoformat(str(txt))
    except ValueError:
        raise HTTPException(400, "data inválida: %s" % txt)


def _periodo(de, ate):
    """[inicio, fim) em horario de Brasilia; padrao: os ultimos 30 dias."""
    hoje = datetime.now(FUSO).date()
    d1 = _dia(ate, hoje)
    d0 = _dia(de, d1 - timedelta(days=29))
    if d0 > d1:
        raise HTTPException(400, "o início do período é depois do fim")
    if (d1 - d0).days > 3660:
        raise HTTPException(400, "período de no máximo 10 anos")
    t0 = datetime(d0.year, d0.month, d0.day, tzinfo=FUSO)
    t1 = datetime(d1.year, d1.month, d1.day, tzinfo=FUSO) + timedelta(days=1)
    return d0, d1, t0, t1


def _uuid_ou_none(s, campo):
    if not s:
        return None
    try:
        return str(uuid.UUID(str(s)))
    except ValueError:
        raise HTTPException(400, "%s inválido" % campo)


def _empresas(cur, u, pedida):
    """(ids alvo, [{id, nome}] que a pessoa pode escolher, id unico ou None).

    O ADMINISTRADOR NAO ESCOLHE: e sempre a empresa dele, qualquer que seja o
    parametro. O suporte escolhe uma ou fica com todas."""
    cur.execute("select id::text, coalesce(name, '') from core.tb_empresas order by name")
    visiveis = [{"id": i, "nome": n} for i, n in cur.fetchall()]
    if u.nivel != "root":
        if not u.id_empresa:
            raise HTTPException(403, "usuário sem empresa")
        return [u.id_empresa], [e for e in visiveis if e["id"] == u.id_empresa], u.id_empresa
    pedida = _uuid_ou_none(pedida, "empresa")
    if pedida:
        if pedida not in {e["id"] for e in visiveis}:
            raise HTTPException(404, "empresa não encontrada")
        return [pedida], visiveis, pedida
    return [e["id"] for e in visiveis], visiveis, None


def _iso(t):
    return t.isoformat(timespec="seconds") if t else None


def _segundos(em, aberta, lote):
    if lote or not aberta or not em:
        return None
    s = (em - aberta).total_seconds()
    return s if s >= 0 else None


def _cadastro(cur, emps, ligs, titular=False):
    """{(empresa, ligacao): (cidade, bairro, endereco, qualificacao, titular)} pela chave do cadastro."""
    nums = sorted({int(l) for l in ligs if str(l).isdigit()})
    if not nums:
        return {}
    cur.execute("""select id_empresa::text, num_ligacao::text, coalesce(cidade, ''), coalesce(nom_bairro, ''),
                          coalesce(end_ligacao, ''), coalesce(qualificacao, ''), """
                + ("coalesce(nom_cliente, '')" if titular else "''") + """
                     from resources_root.cadastro_corsan
                    where id_empresa = any(%s::uuid[]) and num_ligacao = any(%s::bigint[])""",
                (emps, nums))
    return {(e, l): r for e, l, *r in cur.fetchall()}


def _vigentes(cur, emps, ligs):
    """{(empresa, ligacao): (id, acao, quem_nome, em)} — a decisao mais recente de cada ligacao."""
    if not ligs:
        return {}
    cur.execute("""select distinct on (id_empresa, ligacao) id_empresa::text, ligacao, id, acao, quem_nome, em
                     from radar_comercial.seek_decisao
                    where id_empresa = any(%s::uuid[]) and ligacao = any(%s::text[])
                    order by id_empresa, ligacao, em desc, id desc""", (emps, sorted(set(ligs))))
    return {(e, l): (i, a, q, t) for e, l, i, a, q, t in cur.fetchall()}


def _cidade_ok(cidade, cad, chave):
    if not cidade:
        return True
    r = cad.get(chave)
    return bool(r) and r[0].strip().upper() == cidade.strip().upper()


# O TESTE DE DESENVOLVIMENTO FICA FORA DA GESTAO DE PRODUCAO (16/09/2026, docs/PLANO_CIDADE_NOVA.md). Dev e
# producao usam o mesmo banco; as decisoes e aberturas das ligacoes cujo veredito saiu de um processo de dev
# (`percepcao.ambiente = 'desenvolvimento'`) nao entram nos numeros. A API de dev mostra tudo. O indice parcial
# da migracao 0119 guarda so essas linhas, e a consulta nao varre os vereditos.
EM_DEV = os.environ.get("RADAR_AMBIENTE", "").strip() == "desenvolvimento"


def _de_teste(cur, emps):
    """{(empresa, ligacao)} julgadas por processo de desenvolvimento; vazio na API de dev."""
    if EM_DEV:
        return set()
    cur.execute("""select id_empresa::text, ligacao from radar_comercial.ligacao_veredito
                    where id_empresa = any(%s::uuid[]) and percepcao->>'ambiente' = 'desenvolvimento'""", (emps,))
    return set(cur.fetchall())


# ───────────────────────────────────────────────────────── resumo ──
@router.get("/api/seek/gestao/resumo")
def gestao_resumo(de: str = "", ate: str = "", cidade: str = "", empresa: str = "",
                  u: _auth.Usuario = Depends(_quem)):
    """Por usuario: vigentes e total de decisoes por status, taxa de aprovacao, tempo
    medio e mediano entre abrir e decidir, aberturas e ultima atividade."""
    d0, d1, t0, t1 = _periodo(de, ate)
    cidade = (cidade or "").strip()
    t_ini = time.time()
    con = _auth.conectar_como(u)
    try:
        cur = con.cursor()
        emps, visiveis, unica = _empresas(cur, u, empresa)
        cur.execute("""select id, id_empresa::text, ligacao, acao, quem::text, quem_nome, em, aberta_em,
                              lote is not null
                         from radar_comercial.seek_decisao
                        where id_empresa = any(%s::uuid[]) and em >= %s and em < %s""", (emps, t0, t1))
        decisoes = cur.fetchall()
        # O QUE FOI DECIDIDO DEPOIS DO FIM DO PERIODO tira a vigencia de quem veio antes.
        cur.execute("""select distinct id_empresa::text, ligacao from radar_comercial.seek_decisao
                        where id_empresa = any(%s::uuid[]) and em >= %s""", (emps, t1))
        depois = set(cur.fetchall())
        cur.execute("""select id_empresa::text, ligacao, quem::text, quem_nome, aberta_em
                         from radar_comercial.seek_abertura
                        where id_empresa = any(%s::uuid[]) and aberta_em >= %s and aberta_em < %s""",
                    (emps, t0, t1))
        aberturas = cur.fetchall()
        teste = _de_teste(cur, emps)
        ligs = {r[2] for r in decisoes} | {r[1] for r in aberturas}
        cad = _cadastro(cur, emps, ligs)
        quem_ids = sorted({r[4] for r in decisoes if r[4]} | {r[2] for r in aberturas})
        cur.execute("""select us.id::text, us.name, us.email, coalesce(n.codigo, 'user'), us.ativo,
                              us.id_empresa::text, coalesce(n.hierarquia, 0)
                         from core.tb_users us
                         left join core.tb_niveis_user n on n.id = us.id_nivel_user
                        where us.id_empresa = any(%s::uuid[]) or us.id = any(%s::uuid[])""", (emps, quem_ids))
        pessoas = {r[0]: r for r in cur.fetchall()}
    finally:
        con.close()

    cidades = sorted({r[0].strip().upper() for k, r in cad.items() if r[0].strip() and k not in teste})
    decisoes = [r for r in decisoes if (r[1], r[2]) not in teste and _cidade_ok(cidade, cad, (r[1], r[2]))]
    aberturas = [r for r in aberturas if (r[0], r[1]) not in teste and _cidade_ok(cidade, cad, (r[0], r[1]))]

    ultima = {}
    for r in decisoes:
        k = (r[1], r[2])
        if k in depois:
            continue
        if k not in ultima or (r[6], r[0]) > (ultima[k][6], ultima[k][0]):
            ultima[k] = r
    vig_ids = {r[0] for r in ultima.values()}

    def vazio():
        return {"vigentes": dict.fromkeys(ACOES, 0), "decisoes": dict.fromkeys(ACOES, 0), "em_lote": 0,
                "tempos": [], "aberturas": 0, "abertas": set(), "ultima_decisao": None,
                "ultima_abertura": None, "nome": None}

    por = {}
    total = vazio()
    for i, emp, lig, acao, quem, nome, em, aberta, lote in decisoes:
        for alvo in (por.setdefault(quem or "?", vazio()), total):
            alvo["nome"] = alvo["nome"] or nome
            alvo["decisoes"][acao] = alvo["decisoes"].get(acao, 0) + 1
            if i in vig_ids:
                alvo["vigentes"][acao] = alvo["vigentes"].get(acao, 0) + 1
            if lote:
                alvo["em_lote"] += 1
            s = _segundos(em, aberta, lote)
            if s is not None:
                alvo["tempos"].append(s)
            if not alvo["ultima_decisao"] or em > alvo["ultima_decisao"]:
                alvo["ultima_decisao"] = em
    for emp, lig, quem, nome, aberta in aberturas:
        for alvo in (por.setdefault(quem, vazio()), total):
            alvo["nome"] = alvo["nome"] or nome
            alvo["aberturas"] += 1
            alvo["abertas"].add((emp, lig))
            if not alvo["ultima_abertura"] or aberta > alvo["ultima_abertura"]:
                alvo["ultima_abertura"] = aberta
    # QUEM PODE DECIDIR E NAO DECIDIU TAMBEM APARECE, zerado: sumir da lista
    # esconderia justamente quem esta parado.
    for pid, p in pessoas.items():
        if p[5] in emps and p[4] and p[6] >= 40 and pid not in por:
            por[pid] = vazio()

    def saida(pid, a):
        p = pessoas.get(pid)
        vt = sum(a["vigentes"].values())
        dt = sum(a["decisoes"].values())
        tempos = a["tempos"]
        ultima_at = max([t for t in (a["ultima_decisao"], a["ultima_abertura"]) if t], default=None)
        return {"id": pid if pid != "?" else None,
                "nome": (p[1] if p else None) or a["nome"] or "(sem nome)",
                "email": p[2] if p else None, "nivel": p[3] if p else None,
                "ativo": p[4] if p else None, "id_empresa": p[5] if p else None,
                "vigentes": a["vigentes"], "vigentes_total": vt,
                "decisoes": a["decisoes"], "decisoes_total": dt, "em_lote": a["em_lote"],
                "taxa_aprovacao": round(a["vigentes"]["aprovar"] / vt, 4) if vt else None,
                "tempo_medio_s": round(statistics.fmean(tempos), 1) if tempos else None,
                "tempo_mediano_s": round(statistics.median(tempos), 1) if tempos else None,
                "medidas": len(tempos), "aberturas": a["aberturas"], "ligacoes_abertas": len(a["abertas"]),
                "ultima_decisao": _iso(a["ultima_decisao"]), "ultima_abertura": _iso(a["ultima_abertura"]),
                "ultima_atividade": _iso(ultima_at)}

    usuarios = sorted((saida(pid, a) for pid, a in por.items()),
                      key=lambda x: (-(x["decisoes_total"] or 0), -(x["aberturas"] or 0), x["nome"].lower()))
    tot = saida("?", total)
    return {"periodo": {"de": d0.isoformat(), "ate": d1.isoformat()}, "cidade": cidade or None,
            "empresa": unica, "empresas": visiveis if u.nivel == "root" else [], "cidades": cidades,
            "totais": {k: tot[k] for k in ("vigentes", "vigentes_total", "decisoes", "decisoes_total", "em_lote",
                                           "taxa_aprovacao", "tempo_medio_s", "tempo_mediano_s", "medidas",
                                           "aberturas", "ligacoes_abertas", "ultima_atividade")},
            "usuarios": usuarios, "ativos": sum(1 for x in usuarios if x["decisoes_total"] or x["aberturas"]),
            "gerado_em_s": round(time.time() - t_ini, 2)}


# ─────────────────────────────────────────────── ligacoes do usuario ──
@router.get("/api/seek/gestao/ligacoes")
def gestao_ligacoes(quem: str = "", de: str = "", ate: str = "", cidade: str = "", empresa: str = "",
                    acao: str = "", vigentes: int = 0, limite: int = 200, cursor: str = "",
                    u: _auth.Usuario = Depends(_quem)):
    """As decisoes do periodo (de uma pessoa, ou de todas), da mais recente para tras,
    com o status vigente de cada ligacao. `cursor` continua de onde a pagina parou."""
    _, _, t0, t1 = _periodo(de, ate)
    quem_id = _uuid_ou_none(quem, "usuário")
    acoes = [a for a in (acao or "").split(",") if a]
    if any(a not in ACOES for a in acoes):
        raise HTTPException(400, "ação inválida")
    limite = max(1, min(int(limite or 200), 1000))
    cidade = (cidade or "").strip()
    pos = None
    if cursor:
        try:
            em_txt, id_txt = cursor.split("|", 1)
            pos = (datetime.fromisoformat(em_txt), int(id_txt))
        except ValueError:
            raise HTTPException(400, "cursor inválido")
    con = _auth.conectar_como(u)
    try:
        cur = con.cursor()
        emps, _, _ = _empresas(cur, u, empresa)
        teste = _de_teste(cur, emps)
        linhas, varridas, fim = [], 0, False
        while len(linhas) < limite and varridas < MAX_VARRIDAS:
            lote = max(limite * 2, 200)
            cur.execute("""select id, id_empresa::text, ligacao, acao, motivo, quem::text, quem_nome, em,
                                  aberta_em, lote::text
                             from radar_comercial.seek_decisao
                            where id_empresa = any(%s::uuid[]) and em >= %s and em < %s
                              and (%s::uuid is null or quem = %s::uuid)
                              and (cardinality(%s::text[]) = 0 or acao = any(%s::text[]))
                              and (%s::timestamptz is null or (em, id) < (%s::timestamptz, %s::bigint))
                            order by em desc, id desc
                            limit %s""",
                        (emps, t0, t1, quem_id, quem_id, acoes, acoes,
                         pos[0] if pos else None, pos[0] if pos else None, pos[1] if pos else None, lote))
            bloco = cur.fetchall()
            if not bloco:
                fim = True
                break
            varridas += len(bloco)
            pos = (bloco[-1][7], bloco[-1][0])
            ligs = [r[2] for r in bloco]
            cad = _cadastro(cur, emps, ligs)
            vig = _vigentes(cur, emps, ligs)
            for r in bloco:
                k = (r[1], r[2])
                if k in teste or not _cidade_ok(cidade, cad, k):
                    continue
                v = vig.get(k)
                eh_vigente = bool(v and v[0] == r[0])
                if vigentes and not eh_vigente:
                    continue
                c = cad.get(k) or ("", "", "", "", "")
                linhas.append({"id": r[0], "id_empresa": r[1], "ligacao": r[2], "acao": r[3], "motivo": r[4],
                               "quem": r[5], "quem_nome": r[6], "em": _iso(r[7]), "aberta_em": _iso(r[8]),
                               "tempo_s": _segundos(r[7], r[8], r[9]), "lote": r[9], "vigente": eh_vigente,
                               "acao_vigente": v[1] if v else None, "vigente_por": v[2] if v else None,
                               "cidade": c[0], "bairro": c[1], "endereco": c[2]})
                if len(linhas) >= limite:
                    pos = (r[7], r[0])
                    break
            if len(bloco) < lote and len(linhas) < limite:
                fim = True
                break
    finally:
        con.close()
    return {"linhas": linhas, "cursor": None if fim else "%s|%s" % (pos[0].isoformat(), pos[1]),
            "varridas": varridas}


# ────────────────────────────────────────────── historico da ligacao ──
@router.get("/api/seek/gestao/ligacao/{ligacao}")
def gestao_ligacao(ligacao: str, empresa: str = "", u: _auth.Usuario = Depends(_quem)):
    """Toda mudanca de status da ligacao, com comentario, quem e quando; e quem abriu."""
    if not ligacao.isdigit():
        raise HTTPException(400, "ligação inválida")
    con = _auth.conectar_como(u)
    try:
        cur = con.cursor()
        emps, _, _ = _empresas(cur, u, empresa)
        if any(lig == ligacao for _, lig in _de_teste(cur, emps)):
            raise HTTPException(404, "ligação não encontrada")
        cur.execute("""select id_empresa::text, veredito,
                              coalesce(percepcao->'resposta'->>'motivo', justificativa), avaliado_em
                         from radar_comercial.ligacao_veredito
                        where id_empresa = any(%s::uuid[]) and ligacao = %s""", (emps, ligacao))
        v = cur.fetchone()
        emp = v[0] if v else None
        alvo = [emp] if emp else emps
        cad = _cadastro(cur, alvo, [ligacao], titular=False)
        c = next(iter(cad.values()), None)
        cur.execute("""select id, id_empresa::text, acao, motivo, quem_nome, em, aberta_em, lote::text
                         from radar_comercial.seek_decisao
                        where id_empresa = any(%s::uuid[]) and ligacao = %s
                        order by em desc, id desc limit 500""", (alvo, ligacao))
        decisoes = [{"id": i, "id_empresa": e, "acao": a, "motivo": m, "quem_nome": q, "em": _iso(t),
                     "aberta_em": _iso(ab), "tempo_s": _segundos(t, ab, lo), "lote": lo}
                    for i, e, a, m, q, t, ab, lo in cur.fetchall()]
        cur.execute("""select quem_nome, aberta_em from radar_comercial.seek_abertura
                        where id_empresa = any(%s::uuid[]) and ligacao = %s
                        order by aberta_em desc limit 200""", (alvo, ligacao))
        aberturas = [{"quem_nome": q, "aberta_em": _iso(t)} for q, t in cur.fetchall()]
        cur.execute("""select quem_nome, desde from radar_comercial.seek_trava
                        where id_empresa = any(%s::uuid[]) and ligacao = %s and expira_em > now()""",
                    (alvo, ligacao))
        t = cur.fetchone()
    finally:
        con.close()
    if not v and not decisoes:
        raise HTTPException(404, "ligação não encontrada")
    return {"ligacao": ligacao, "id_empresa": emp,
            "base": {"cidade": c[0], "bairro": c[1], "endereco": c[2], "qualificacao": c[3]} if c else None,
            "ia": {"veredito": v[1], "motivo": v[2], "avaliado_em": _iso(v[3])} if v else None,
            "decisoes": decisoes, "aberturas": aberturas,
            "trava": {"quem_nome": t[0], "desde": _iso(t[1])} if t else None}


# ────────────────────────────────────────────────────── exportacao ──
class ExportarEntrada(BaseModel):
    formato: str = "xlsx"
    de: str | None = None
    ate: str | None = None
    empresa: str | None = None
    quem: str | None = None
    acoes: list[str] = []
    cidade: str | None = None
    vigentes: bool = False
    colunas: list[str] = list(COLUNAS_PADRAO)


def _ip(request: Request):
    for cab in ("x-forwarded-for", "x-real-ip"):
        v = (request.headers.get(cab) or "").split(",")[0].strip()
        if v:
            return v[:64]
    return request.client.host if request.client else None


def _tempo_txt(s):
    if s is None:
        return None
    s = int(round(s))
    return "%d:%02d:%02d" % (s // 3600, (s % 3600) // 60, s % 60)


def _local(t):
    return t.astimezone(FUSO).strftime("%d/%m/%Y %H:%M:%S") if t else None


#: O que faz o Excel ler a celula como formula (injecao de formula em CSV/XLSX).
_INICIO_FORMULA = ("=", "+", "-", "@", "\t", "\r")


def _sem_formula_csv(x):
    if isinstance(x, str) and x[:1] in _INICIO_FORMULA:
        return "'" + x
    return x


@router.post("/api/seek/gestao/exportar")
def gestao_exportar(e: ExportarEntrada, request: Request, u: _auth.Usuario = Depends(_quem)):
    """Gera o arquivo com as decisoes filtradas, registra a exportacao e entrega.

    O ARQUIVO NASCE INTEIRO ANTES DE SAIR: ler o banco a velocidade do banco e
    soltar a conexao, e so entao mandar. Entregar direto do cursor prenderia uma
    conexao do pooler pelo tempo do download de quem baixa devagar."""
    formato = (e.formato or "").lower()
    if formato not in ("xlsx", "csv"):
        raise HTTPException(400, "formato: xlsx ou csv")
    colunas = [c for c in (e.colunas or []) if c in NOMES_COLUNAS]
    if not colunas:
        raise HTTPException(400, "escolha ao menos uma coluna")
    acoes = [a for a in (e.acoes or []) if a in ACOES]
    d0, d1, t0, t1 = _periodo(e.de, e.ate)
    quem_id = _uuid_ou_none(e.quem, "usuário")
    cidade = (e.cidade or "").strip()
    precisa_cad = cidade or any(c in colunas for c in ("cidade", "bairro", "endereco", "titular", "qualificacao"))
    precisa_ia = any(c in colunas for c in ("veredito_ia", "motivo_ia", "julgada_em"))
    precisa_vig = e.vigentes or any(c in colunas for c in ("vigente", "decisao_vigente"))

    # ARQUIVO DE VERDADE, e nao SpooledTemporaryFile: no Python 3.10 da imagem ele
    # nao tem `readable()`/`seekable()` e o TextIOWrapper do CSV recusa. O /tmp do
    # conteiner e tmpfs (so-leitura no resto), entao continua em memoria.
    tmp = tempfile.TemporaryFile(mode="w+b", dir="/tmp")
    n = 0
    con = _auth.conectar_como(u)
    try:
        cur = con.cursor()
        emps, visiveis, unica = _empresas(cur, u, e.empresa)
        nome_emp = {x["id"]: x["nome"] for x in visiveis}
        teste = _de_teste(cur, emps)
        if formato == "csv":
            texto = io.TextIOWrapper(tmp, encoding="utf-8-sig", newline="")
            escritor = csv.writer(texto, delimiter=";", quoting=csv.QUOTE_MINIMAL)
            escritor.writerow([NOMES_COLUNAS[c] for c in colunas])

            def escrever(valores):
                escritor.writerow([_sem_formula_csv(x) for x in valores])
        else:
            from openpyxl import Workbook
            from openpyxl.cell import WriteOnlyCell
            from openpyxl.styles import Font, PatternFill
            wb = Workbook(write_only=True)
            ws = wb.create_sheet("Decisões")
            cab = []
            for c in colunas:
                cel = WriteOnlyCell(ws, value=NOMES_COLUNAS[c])
                cel.font = Font(bold=True, color="FFFFFF")
                cel.fill = PatternFill("solid", fgColor="0B2E59")
                cab.append(cel)
            ws.append(cab)

            def escrever(valores):
                linha = []
                for x in valores:
                    if isinstance(x, str) and x[:1] in _INICIO_FORMULA:
                        # TEXTO QUE COMECA COM "=" VIRA FORMULA no openpyxl: o
                        # comentario escrito por um editor seria executado no
                        # Excel de quem exporta. Forca a celula como texto.
                        cel = WriteOnlyCell(ws, value=x)
                        cel.data_type = "s"
                        linha.append(cel)
                    else:
                        linha.append(x)
                ws.append(linha)

        # UM CURSOR DO LADO DO SERVIDOR, lido em blocos: a lista inteira nao sobe
        # de uma vez para a memoria da API.
        sc = con.cursor(name="seek_exportar_%s" % uuid.uuid4().hex[:8])
        sc.itersize = 5000
        sc.execute("""select id, id_empresa::text, ligacao, acao, motivo, quem_nome, em, aberta_em, lote is not null
                        from radar_comercial.seek_decisao
                       where id_empresa = any(%s::uuid[]) and em >= %s and em < %s
                         and (%s::uuid is null or quem = %s::uuid)
                         and (cardinality(%s::text[]) = 0 or acao = any(%s::text[]))
                       order by em, id""", (emps, t0, t1, quem_id, quem_id, acoes, acoes))
        excedeu = False
        while not excedeu:
            bloco = sc.fetchmany(5000)
            if not bloco:
                break
            ligs = [r[2] for r in bloco]
            cad = _cadastro(cur, emps, ligs, titular="titular" in colunas) if precisa_cad else {}
            vig = _vigentes(cur, emps, ligs) if precisa_vig else {}
            ia = {}
            if precisa_ia:
                cur.execute("""select id_empresa::text, ligacao, veredito,
                                      coalesce(percepcao->'resposta'->>'motivo', justificativa), avaliado_em
                                 from radar_comercial.ligacao_veredito
                                where id_empresa = any(%s::uuid[]) and ligacao = any(%s::text[])""",
                            (emps, sorted(set(ligs))))
                ia = {(a, b): (c, d, f) for a, b, c, d, f in cur.fetchall()}
            for i, emp, lig, acao, motivo, quem_nome, em, aberta, lote in bloco:
                k = (emp, lig)
                if k in teste or not _cidade_ok(cidade, cad, k):
                    continue
                v = vig.get(k)
                eh_vig = bool(v and v[0] == i)
                if e.vigentes and not eh_vig:
                    continue
                n += 1
                if n > MAX_LINHAS:
                    excedeu = True
                    break
                c = cad.get(k) or ("", "", "", "", "")
                j = ia.get(k) or (None, None, None)
                s = _segundos(em, aberta, lote)
                valores = {
                    "ligacao": lig, "empresa": nome_emp.get(emp, emp), "cidade": c[0], "bairro": c[1],
                    "endereco": c[2], "qualificacao": (c[3] or "").replace("_", " "), "titular": c[4],
                    "veredito_ia": ROTULO_IA.get(j[0], j[0]), "motivo_ia": j[1], "julgada_em": _local(j[2]),
                    "decisao": ROTULO_ACAO.get(acao, acao), "comentario": motivo, "quem": quem_nome,
                    "quando": _local(em), "aberta_em": _local(aberta), "tempo_avaliacao": _tempo_txt(s),
                    "tempo_avaliacao_s": int(round(s)) if s is not None else None,
                    "em_lote": "sim" if lote else "não", "vigente": "sim" if eh_vig else "não",
                    "decisao_vigente": ROTULO_ACAO.get(v[1], v[1]) if v else None,
                }
                escrever([valores[col] for col in colunas])
        sc.close()
        if excedeu:
            con.rollback()
            raise HTTPException(413, "mais de %s linhas: refine o período ou os filtros" % format(MAX_LINHAS, ",d").replace(",", "."))

        if formato == "csv":
            texto.flush()
            texto.detach()
        else:
            wf = wb.create_sheet("Filtros")
            for par in (("Exportado por", u.nome or u.email), ("Em", _local(datetime.now(FUSO))),
                        ("Período", "%s a %s" % (d0.strftime("%d/%m/%Y"), d1.strftime("%d/%m/%Y"))),
                        ("Empresa", nome_emp.get(unica, unica) if unica else "todas"),
                        ("Usuário", e.quem or "todos"), ("Status", ", ".join(acoes) or "todos"),
                        ("Cidade", cidade or "todas"), ("Só decisões vigentes", "sim" if e.vigentes else "não"),
                        ("Linhas", n)):
                wf.append(list(par))
            wb.save(tmp)
        tmp.seek(0)

        filtros = {"de": d0.isoformat(), "ate": d1.isoformat(), "empresa": unica, "quem": quem_id,
                   "acoes": acoes, "cidade": cidade or None, "vigentes": bool(e.vigentes)}
        cur.execute("""insert into radar_comercial.seek_exportacao
                           (id_empresa, quem, quem_nome, formato, filtros, colunas, linhas, ip, agente)
                       values (%s, %s, %s, %s, %s::jsonb, %s, %s, %s, %s) returning id""",
                    (unica, u.id, u.nome or u.email, formato, json.dumps(filtros), colunas, n,
                     _ip(request), (request.headers.get("user-agent") or "")[:300]))
        reg = cur.fetchone()[0]
        con.commit()
    except Exception:
        tmp.close()
        raise
    finally:
        con.close()

    nome = "seek_decisoes_%s.%s" % (datetime.now(FUSO).strftime("%Y%m%d_%H%M"), formato)
    tipo = ("text/csv; charset=utf-8" if formato == "csv"
            else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def pedacos():
        try:
            while True:
                b = tmp.read(256 * 1024)
                if not b:
                    break
                yield b
        finally:
            tmp.close()

    return StreamingResponse(pedacos(), media_type=tipo, headers={
        "Content-Disposition": 'attachment; filename="%s"' % nome, "Cache-Control": "no-store",
        "X-Exportacao-Id": str(reg), "X-Exportacao-Linhas": str(n)})


@router.get("/api/seek/gestao/exportacoes")
def gestao_exportacoes(empresa: str = "", limite: int = 100, u: _auth.Usuario = Depends(_quem)):
    limite = max(1, min(int(limite or 100), 500))
    con = _auth.conectar_como(u)
    try:
        cur = con.cursor()
        emps, visiveis, unica = _empresas(cur, u, empresa)
        nome_emp = {x["id"]: x["nome"] for x in visiveis}
        cur.execute("""select id, id_empresa::text, quem_nome, em, formato, filtros, colunas, linhas, ip
                         from radar_comercial.seek_exportacao
                        where id_empresa = any(%s::uuid[]) or (%s and id_empresa is null)
                        order by em desc limit %s""", (emps, u.nivel == "root" and not unica, limite))
        return {"exportacoes": [
            {"id": i, "id_empresa": e, "empresa": nome_emp.get(e, "todas" if e is None else e), "quem_nome": q,
             "em": _iso(t), "formato": f, "filtros": fi, "colunas": co, "linhas": n, "ip": ip}
            for i, e, q, t, f, fi, co, n, ip in cur.fetchall()],
            "colunas": [{"id": c, "nome": nm, "padrao": c in COLUNAS_PADRAO} for c, nm in COLUNAS]}
    finally:
        con.close()


# ────────────────────────────────────────────────────────── a pagina ──
@router.get("/gestao")
def gestao_pagina():
    """A pagina e publica como a `/` da SEEK: o `sessao.js` pede o login, e os dados
    so vem pelas rotas acima, que exigem administrador."""
    alvo = FRONT / "gestao.html"
    if not alvo.exists():
        raise HTTPException(404, "gestao.html não gerada — rode python3 frontend/seek_fonte/montar.py")
    html = alvo.read_text(encoding="utf-8")
    js = FRONT / "sessao.js"
    if js.exists():
        html = html.replace("/static/sessao.js", "/static/sessao.js?v=%d" % int(js.stat().st_mtime))
    return Response(html, media_type="text/html", headers={"Cache-Control": "no-cache"})
