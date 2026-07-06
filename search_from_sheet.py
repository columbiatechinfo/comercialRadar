"""
search_from_sheet.py — Fluxo NOVO (a parte do pipeline principal)

Entrada: uma planilha (.xlsx/.csv) com POIs que você já sabe que existem ou que
possivelmente existam. Colunas aceitas (detecção flexível por cabeçalho):
  - nome      (obrigatório)  — "nome", "nome do local", "estabelecimento", "name"
  - endereco  (opcional)     — "endereco", "endereço", "address", "logradouro"
  - lat/lng   (opcional)     — "lat"/"latitude" e "lng"/"lon"/"longitude",
                                ou uma coluna única "coordenada" no formato "lat,lng"

Para cada linha, busca no Google Maps com a MESMA metodologia v2 (fill robusto,
proxy estático, fingerprint, humanização) e coleta dados completos:
  básicos + horários + até MAX_FOTOS fotos + até MAX_REVIEWS avaliações.

Saída: <planilha>_db.json (formato normalizado p/ o ingester Prisma).
Depois: npx ts-node src/ingest.ts <planilha>_db.json  (ou rode com --ingest).

USO:
  py search_from_sheet.py minha_planilha.xlsx [--workers 6] [--no-proxy] [--cidade "Canoas RS"]
"""

import re
import json
import time
import asyncio
import argparse
import shutil
import subprocess
from pathlib import Path
from urllib.parse import unquote

from playwright.async_api import async_playwright

import config
from proxy_pool import ProxyPool
from human_browser import HumanSession
from search_pois_v2 import (
    haversine, similaridade, salvar_json,
    aguarda_painel_ou_lista, extract_panel, selecionar_melhor_card,
    _preencher_busca, abrir_maps, abrir_coordenada_para_n2,
)
from extract_full import enriquecer_poi
import area_utils


# ──────────────────────────────────────────────────────────────────────────
# Leitura da planilha
# ──────────────────────────────────────────────────────────────────────────
_ALIASES = {
    "nome": ["nome", "nome do local", "estabelecimento", "local", "name", "razao social", "razão social"],
    # endereço completo tem prioridade sobre logradouro (rua só)
    "endereco": ["endereco_completo", "endereço_completo", "endereco completo", "endereço completo",
                 "endereco", "endereço", "address", "logradouro"],
    "lat": ["lat", "latitude"],
    "lng": ["lon", "lng", "long", "longitude"],
    "coordenada": ["coordenada", "coordenadas", "coord", "latlng", "lat,lng"],
    "uf": ["uf", "estado", "sigla_uf", "sigla"],
}

_UFS = {"AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA", "MT", "MS",
        "MG", "PA", "PB", "PR", "PE", "PI", "RJ", "RN", "RS", "RO", "RR", "SC",
        "SP", "SE", "TO"}


def uf_do_endereco(endereco: str):
    """Extrai a sigla da UF de um endereço do Google Maps (ex.: '... - PI, 64200-000')."""
    if not endereco:
        return None
    # UF antes do CEP: '... PI, 64200-000' ou '... - PI'
    for m in re.finditer(r"\b([A-Z]{2})\b", endereco):
        if m.group(1) in _UFS:
            uf = m.group(1)
    # pega a última ocorrência válida (a UF vem no fim do endereço)
    achados = [g for g in re.findall(r"\b([A-Z]{2})\b", endereco) if g in _UFS]
    return achados[-1] if achados else None


def _achar_col(headers, chave):
    hl = [(_norm(h) if h else "") for h in headers]
    for alvo in _ALIASES[chave]:
        for i, h in enumerate(hl):
            if h == alvo:
                return i
    # match parcial
    for alvo in _ALIASES[chave]:
        for i, h in enumerate(hl):
            if alvo in h:
                return i
    return None


def _norm(s):
    return str(s).strip().lower()


# ──────────────────────────────────────────────────────────────────────────
# Match de nome por CONTINÊNCIA (nome conhecido x nome do Maps)
# ──────────────────────────────────────────────────────────────────────────
import unicodedata

_STOP = {"de", "da", "do", "dos", "das", "e", "o", "a", "&", "-", "|", "",
         "em", "no", "na", "ltda", "me", "epp"}


def _norm_nome(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-zA-Z0-9\s]", " ", s.lower())
    return re.sub(r"\s+", " ", s).strip()


def _tokens(s: str) -> set:
    return {t for t in _norm_nome(s).split() if t not in _STOP and len(t) > 1}


def nome_match(a: str, b: str, dist) -> tuple[bool, float]:
    """
    Decide se o nome da planilha (a) casa com o nome do Maps (b), usando:
      1. similaridade global (SequenceMatcher) — nomes quase iguais
      2. CONTINÊNCIA de tokens — 'Eletrônica Agamenon' ⊂ 'Eletrônica Agamenon | Instrumentos'
      3. substring direto
    Continência exige proximidade (dist) para não casar homônimos distantes.
    Retorna (match, score_similaridade_global).
    """
    na, nb = _norm_nome(a), _norm_nome(b)
    ra = similaridade(na, nb) if na and nb else 0.0

    # PORTÃO DE DISTÂNCIA: com coordenada da planilha, um match real está PERTO.
    # Sem isso, nomes iguais a lugares famosos (Tóquio, K2, São Francisco) casavam
    # com o lugar famoso a milhares de km. > 20km da coordenada = outro lugar.
    if dist is not None and dist > 20000:
        return False, ra

    if ra >= 0.72:
        return True, ra

    ta, tb = _tokens(a), _tokens(b)
    if ta and tb:
        cont = len(ta & tb) / min(len(ta), len(tb))
        if cont >= 0.6 and dist is not None and dist <= 400:
            return True, ra
        if cont >= 0.85 and dist is not None and dist <= 1500:
            return True, ra

    if na and nb and (na in nb or nb in na) and min(len(na), len(nb)) >= 5:
        if dist is not None and dist <= 400:
            return True, ra

    return False, ra


def _parse_coord(txt):
    if not txt:
        return None, None
    m = re.search(r"(-?\d+\.\d+)\s*[,;]\s*(-?\d+\.\d+)", str(txt))
    if m:
        return float(m.group(1)), float(m.group(2))
    return None, None


def ler_planilha(path: Path) -> list:
    linhas = []
    if path.suffix.lower() == ".csv":
        import csv
        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]

    if not rows:
        return []

    headers = rows[0]
    ci_nome = _achar_col(headers, "nome")
    ci_end = _achar_col(headers, "endereco")
    ci_lat = _achar_col(headers, "lat")
    ci_lng = _achar_col(headers, "lng")
    ci_coord = _achar_col(headers, "coordenada")
    ci_uf = _achar_col(headers, "uf")

    if ci_nome is None:
        raise ValueError("Coluna de nome não encontrada na planilha (cabeçalhos: %s)" % headers)

    for i, r in enumerate(rows[1:]):
        if not r or ci_nome >= len(r) or not r[ci_nome]:
            continue
        nome = str(r[ci_nome]).strip()
        endereco = str(r[ci_end]).strip() if (ci_end is not None and ci_end < len(r) and r[ci_end]) else ""
        lat = lng = None
        if ci_lat is not None and ci_lng is not None:
            try:
                lat = float(str(r[ci_lat]).replace(",", ".")) if r[ci_lat] not in (None, "") else None
                lng = float(str(r[ci_lng]).replace(",", ".")) if r[ci_lng] not in (None, "") else None
            except Exception:
                pass
        if (lat is None or lng is None) and ci_coord is not None and ci_coord < len(r):
            lat, lng = _parse_coord(r[ci_coord])
        uf = None
        if ci_uf is not None and ci_uf < len(r) and r[ci_uf]:
            uf = str(r[ci_uf]).strip().upper()[:2]
        linhas.append({"_row": i, "nome": nome, "endereco": endereco, "lat": lat, "lng": lng, "uf": uf})
    return linhas


# ──────────────────────────────────────────────────────────────────────────
# Busca de uma linha
# ──────────────────────────────────────────────────────────────────────────
def extrair_place_id(maps_url: str):
    if not maps_url:
        return None
    m = re.search(r"!16s([^!?&]+)", maps_url)
    if m:
        return unquote(m.group(1)).lstrip("/")
    m = re.search(r"!1s(0x[0-9a-f]+:0x[0-9a-f]+)", maps_url)
    return m.group(1) if m else None


async def coletar_candidatos_proximos(sess, lat, lng, nome, max_c: int = 6) -> list:
    """
    Busca "próximo daqui": posiciona na coordenada, clica em 'Próximo' (pesquisar
    nas proximidades), busca o nome e coleta os estabelecimentos vizinhos
    retornados (nome, endereço, coord, place_id, distância).
    """
    page = sess.page
    cands = []
    if lat is None or lng is None:
        return cands
    try:
        if await abrir_coordenada_para_n2(page, lat, lng) != "ok":
            return cands
        prox = page.get_by_role("button", name="Próximo", exact=True)
        try:
            await prox.first.wait_for(state="visible", timeout=config.WAIT_PROXIMO_MS)
        except Exception:
            prox = page.locator('[aria-label="Próximo"]')
            try:
                await prox.first.wait_for(state="visible", timeout=2500)
            except Exception:
                return cands
        await prox.first.click()
        await page.wait_for_timeout(800)
        await _preencher_busca(page, nome)
        await page.keyboard.press("Enter")
        await page.wait_for_timeout(config.WAIT_APOS_BUSCA_MS)

        MAX_VIZ = 20000  # candidato vizinho só vale se estiver ≤ 20km da coordenada
        res = await aguarda_painel_ou_lista(page)
        if res == "painel":
            poi = await extract_panel(page)
            if poi.get("nome"):
                d = None
                if poi.get("maps_lat") and poi.get("maps_lng"):
                    d = haversine(lat, lng, poi["maps_lat"], poi["maps_lng"])
                if d is not None and d <= MAX_VIZ:
                    cands.append({
                        "nome": poi["nome"], "endereco": poi.get("endereco", ""),
                        "lat": poi.get("maps_lat"), "lng": poi.get("maps_lng"),
                        "place_id": extrair_place_id(poi.get("maps_url", "")),
                        "maps_url": poi.get("maps_url", ""),
                        "dist_m": round(d, 1),
                    })
        elif res == "lista":
            cards = page.locator('div[role="feed"] div[role="article"], div.Nv2PK')
            n = await cards.count()
            for i in range(min(n, max_c)):
                card = cards.nth(i)
                try:
                    nm = ""
                    for ns in ["div.qBF1Pd", "span.fontHeadlineSmall"]:
                        try:
                            nm = await card.locator(ns).first.inner_text(timeout=600)
                            if nm:
                                break
                        except Exception:
                            continue
                    if not nm:
                        continue
                    href = await card.locator('a[href*="/maps/place/"]').first.get_attribute("href", timeout=600)
                    clat = clng = None
                    m = re.search(r"!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)", href or "")
                    if m:
                        clat, clng = float(m.group(1)), float(m.group(2))
                    d = haversine(lat, lng, clat, clng) if clat is not None else None
                    if d is None or d > MAX_VIZ:
                        continue
                    cands.append({
                        "nome": nm.strip(), "endereco": "", "lat": clat, "lng": clng,
                        "place_id": extrair_place_id(href or ""), "maps_url": href or "",
                        "dist_m": round(d, 1),
                    })
                except Exception:
                    continue
    except Exception:
        pass
    return cands


async def _abrir_url_e_extrair(sess, url):
    """Navega para a URL de um lugar e extrai os dados básicos do painel."""
    if not url:
        return None
    try:
        await sess.page.goto(url, wait_until="domcontentloaded", timeout=30000)
        await sess.page.wait_for_timeout(1500)
        if await aguarda_painel_ou_lista(sess.page) != "painel":
            return None
        poi = await extract_panel(sess.page)
        return poi if poi.get("nome") else None
    except Exception:
        return None


async def buscar_linha(sess, item, cidade, sessao_nome, target_uf=None, recuperar=False) -> dict:
    page = sess.page
    nome_in = item["nome"]
    endereco_in = item.get("endereco", "")
    olat, olng = item.get("lat"), item.get("lng")

    # Monta a query: nome + endereço (ou nome + cidade). Já temos o nome certo,
    # então a busca direta é mais robusta que a dança de coordenada.
    if endereco_in:
        query = f"{nome_in}, {endereco_in}"
    else:
        query = f"{nome_in}, {cidade}" if cidade else nome_in

    registro = {
        "fonte": "planilha",
        "sessao": sessao_nome,
        "_row": item.get("_row"),
        "nome_planilha": nome_in,
        "endereco_planilha": endereco_in,
        "lat_origem": olat,
        "lng_origem": olng,
        "status": "nao_encontrado",
        "match_valido": False,
        "fotos": [], "comentarios": [], "horarios": {},
    }

    try:
        url_antes = page.url
        await _preencher_busca(page, query)
        await page.keyboard.press("Enter")

        # Espera o resultado NOVO (a URL muda para /place/ ou /search/), em vez de
        # ler o painel obsoleto do POI anterior que ainda está na tela.
        import time as _t
        fim = _t.time() + config.WAIT_PAINEL_MS / 1000
        while _t.time() < fim:
            u = page.url
            if u != url_antes and ("/place/" in u or "/search/" in u):
                break
            await page.wait_for_timeout(300)
        await page.wait_for_timeout(config.WAIT_APOS_BUSCA_MS)

        res = await aguarda_painel_ou_lista(page)

        poi = None
        if res == "painel":
            poi = await extract_panel(page)
        elif res == "lista":
            poi = await selecionar_melhor_card(sess, nome_in, olat or -29.9, olng or -51.18)

        async def _sucesso(poi_ok, dist_ok, sim_ok, status):
            poi_ok = await enriquecer_poi(sess, poi_ok)
            d = None
            if olat and olng and poi_ok.get("maps_lat") and poi_ok.get("maps_lng"):
                d = haversine(olat, olng, poi_ok["maps_lat"], poi_ok["maps_lng"])
            registro.update({
                "nome": poi_ok.get("nome", ""), "categoria": poi_ok.get("categoria", ""),
                "endereco": poi_ok.get("endereco", ""), "telefone": poi_ok.get("telefone", ""),
                "website": poi_ok.get("website", ""), "avaliacao": poi_ok.get("avaliacao", ""),
                "total_avaliacoes": poi_ok.get("total_avaliacoes", 0), "plus_code": poi_ok.get("plus_code", ""),
                "status_horario": poi_ok.get("status_horario", ""),
                "maps_lat": poi_ok.get("maps_lat"), "maps_lng": poi_ok.get("maps_lng"),
                "maps_url": poi_ok.get("maps_url", ""), "place_id": extrair_place_id(poi_ok.get("maps_url", "")),
                "status": status, "distancia_m": round((dist_ok if dist_ok is not None else d), 1) if (dist_ok is not None or d is not None) else None,
                "similaridade": round(sim_ok, 3) if sim_ok is not None else None, "match_valido": True,
                "fotos": poi_ok.get("fotos", []), "comentarios": poi_ok.get("comentarios", []),
                "horarios": poi_ok.get("horarios", {}),
            })

        async def _recuperar():
            # Busca "próximo daqui": coleta vizinhos, tenta casar por nome entre eles.
            if not (recuperar and olat and olng):
                return
            cands = await coletar_candidatos_proximos(sess, olat, olng, nome_in)
            registro["candidatos_proximos"] = cands
            for c in cands:
                m2, s2 = nome_match(nome_in, c.get("nome", ""), c.get("dist_m"))
                if not m2:
                    continue
                poi2 = await _abrir_url_e_extrair(sess, c.get("maps_url"))
                if not poi2 or not poi2.get("nome"):
                    continue
                uf2 = uf_do_endereco(poi2.get("endereco", ""))
                if target_uf and uf2 and uf2 != target_uf:
                    continue
                await _sucesso(poi2, c.get("dist_m"), s2, "recuperado_proximo")
                return

        if not poi or not poi.get("nome"):
            await _recuperar()
            return registro

        # Decide o match ANTES de enriquecer (não gasta banda em divergente).
        dist = None
        if olat and olng and poi.get("maps_lat") and poi.get("maps_lng"):
            dist = haversine(olat, olng, poi["maps_lat"], poi["maps_lng"])
        match, sim = nome_match(nome_in, poi.get("nome", ""), dist)

        # Filtro de UF: o match tem que estar no mesmo estado do run.
        uf_maps = uf_do_endereco(poi.get("endereco", ""))
        if match and target_uf and uf_maps and uf_maps != target_uf:
            registro.update({"status": "fora_da_uf", "nome": poi.get("nome", ""),
                             "similaridade": round(sim, 3),
                             "distancia_m": round(dist, 1) if dist is not None else None,
                             "uf_encontrada": uf_maps})
            await _recuperar()
            return registro

        if not match:
            registro.update({"status": "encontrado_divergente", "nome": poi.get("nome", ""),
                             "similaridade": round(sim, 3),
                             "distancia_m": round(dist, 1) if dist is not None else None})
            await _recuperar()
            return registro

        await _sucesso(poi, dist, sim, "ok")
    except Exception as e:
        registro["status"] = "erro"
        registro["erro"] = str(e)[:150]

    return registro


# ──────────────────────────────────────────────────────────────────────────
# Worker
# ──────────────────────────────────────────────────────────────────────────
def _chunk(itens, lo, hi):
    import random
    rng = random.Random(7)
    out, i = [], 0
    while i < len(itens):
        if len(itens) - i <= hi:
            out.append(itens[i:]); break
        t = rng.randint(lo, hi)
        if len(itens) - i - t < lo:
            t = len(itens) - i - lo
        out.append(itens[i:i + t]); i += t
    return out


async def worker(wid, queue, state, pw, pool, cidade, sessao_nome, counter, total, out_json, lock, metrics, usar_proxy, target_uf=None, recuperar=False, headless=True, poligono=None):
    while True:
        try:
            bidx, batch = queue.get_nowait()
        except asyncio.QueueEmpty:
            return

        proxy = await pool.acquire_blocking() if usar_proxy else None
        if usar_proxy and not proxy:
            queue.put_nowait((bidx, batch)); await asyncio.sleep(5); continue

        ip_label = f"{proxy['address']}:{proxy['port']}" if proxy else "direto"
        profile = config.BROWSER_PROFILES_DIR / f"sheet_w{wid}_b{bidx}"
        sess = None
        try:
            tz_lng = next((b.get("lng") for b in batch if b.get("lng")), None)
            sess = await HumanSession.create(pw, proxy, profile, layer="maps", tz_hint_lng=tz_lng, headless=headless)
            if not await abrir_maps(sess):
                queue.put_nowait((bidx, batch))
                if proxy: await pool.mark_cooldown(proxy, 600)
                continue

            for j, item in enumerate(batch):
                if await sess.is_captcha():
                    print(f"\n  🚫 [W{wid}] CAPTCHA lote {bidx} IP={ip_label}")
                    break
                rec = await buscar_linha(sess, item, cidade, sessao_nome, target_uf, recuperar)
                area_utils.gate_registro(rec, poligono)
                async with lock:
                    state["results"].append(rec)
                    salvar_json(out_json, state["results"])
                counter["done"] += 1
                metrics["pois"] += 1
                mark = "✅" if rec["match_valido"] else "❌"
                nf, nc = len(rec.get("fotos", [])), len(rec.get("comentarios", []))
                print(f"\r📋 [W{wid}] IP={ip_label} | {mark} {counter['done']}/{total} "
                      f"{rec['status']:20} 📷{nf} 💬{nc} | {item['nome'][:28]}", flush=True)
                if j < len(batch) - 1:
                    await sess.humanized_wait()
        except Exception as e:
            print(f"\n  W{wid} erro lote {bidx}: {str(e)[:90]}")
        finally:
            if sess:
                try:
                    await sess.close(); metrics["bytes"] += sess.bytes_used
                except Exception:
                    pass
            shutil.rmtree(profile, ignore_errors=True)
            if proxy: await pool.release(proxy)


# ──────────────────────────────────────────────────────────────────────────
# Orquestração
# ──────────────────────────────────────────────────────────────────────────
async def run(sheet_path: Path, n_workers: int, cidade: str, usar_proxy: bool, limit: int = 0, retry_failed: bool = False, recuperar: bool = False, headless: bool = True, gemini_direto: bool = False, area_path=None):
    itens = ler_planilha(sheet_path)
    if not itens:
        print("Planilha vazia ou sem coluna de nome.")
        return

    poligono = area_utils.carregar_area(area_path) if area_path else None
    if area_path and not poligono:
        print(f"⚠️  Área inválida/não encontrada: {area_path} (seguindo SEM gate de área)")

    sessao_nome = sheet_path.stem
    out_json = sheet_path.parent / f"{sheet_path.stem}_db.json"

    # UF alvo (uma UF por execução): maioria da coluna UF da planilha
    from collections import Counter
    ufs = Counter(it["uf"] for it in itens if it.get("uf"))
    target_uf = ufs.most_common(1)[0][0] if ufs else None

    # Retomada / retry: define o que já está pronto e o que reprocessar
    all_existing = []
    if out_json.exists():
        try:
            all_existing = json.loads(out_json.read_text(encoding="utf-8"))
        except Exception:
            all_existing = []

    if retry_failed:
        # 'Feito' = válido (ok + recuperados). Falhas são reprocessáveis.
        feitos = {r.get("_row") for r in all_existing if r.get("match_valido") and r.get("_row") is not None}
        if all_existing:
            print(f"\n🔁 Retry-failed: {len(feitos)} válidos preservados, reprocessando falhas.")
    else:
        feitos = {r.get("_row") for r in all_existing if r.get("_row") is not None}
        if all_existing:
            print(f"\n♻️  Retomando: {len(feitos)} linhas já processadas.")

    pendentes = [it for it in itens if it.get("_row") not in feitos]
    # --limit: limita as PENDENTES a reprocessar (útil p/ observar poucos falhos)
    if limit > 0:
        pendentes = pendentes[:limit]

    # Preserva tudo do JSON, EXCETO as linhas que vamos reprocessar agora
    rows_reproc = {it.get("_row") for it in pendentes}
    results_existentes = [r for r in all_existing if r.get("_row") not in rows_reproc]

    print(f"\n📋 ComercialRadar — Busca por planilha")
    print(f"   Planilha : {sheet_path.name}")
    print(f"   POIs     : {len(itens)} | Pendentes: {len(pendentes)} | Proxy: {'sim' if usar_proxy else 'NÃO (direto)'} | Workers: {n_workers}")
    print(f"   Cidade   : {cidade or '(usar endereço da planilha)'} | UF alvo: {target_uf or '(qualquer)'}")
    print(f"   Área     : {'polígono com %d vértices (gate ATIVO)' % len(poligono) if poligono else '(sem gate de área)'}\n")

    if not pendentes:
        print("✅ Todas as linhas já foram processadas.")
        return out_json

    # ── MODO GEMINI-DIRETO: resíduo que já falhou no Maps → direto pro Gemini ──
    if gemini_direto:
        _gemini_direto(pendentes, results_existentes, target_uf, sessao_nome, out_json, poligono)
        return out_json

    lotes = _chunk(pendentes, config.BATCH_MIN, config.BATCH_MAX)
    state = {"results": list(results_existentes)}
    lock = asyncio.Lock()
    counter = {"done": len(feitos)}
    total = len(itens)
    metrics = {"bytes": 0, "pois": 0, "inicio": time.time()}

    pool = ProxyPool().start() if usar_proxy else None
    queue = asyncio.Queue()
    for i, b in enumerate(lotes):
        queue.put_nowait((i, b))

    config.BROWSER_PROFILES_DIR.mkdir(parents=True, exist_ok=True)
    n_ef = min(n_workers, config.MAX_WORKERS, len(lotes))
    async with async_playwright() as pw:
        await asyncio.gather(*[
            worker(i, queue, state, pw, pool, cidade, sessao_nome, counter, total, out_json, lock, metrics, usar_proxy, target_uf, recuperar, headless, poligono)
            for i in range(n_ef)
        ])
    shutil.rmtree(config.BROWSER_PROFILES_DIR, ignore_errors=True)

    dur = time.time() - metrics["inicio"]
    ok = sum(1 for r in state["results"] if r.get("match_valido"))
    mb = metrics["bytes"] / (1024 * 1024)
    tot_fotos = sum(len(r.get("fotos", [])) for r in state["results"])
    tot_rev = sum(len(r.get("comentarios", [])) for r in state["results"])
    print(f"\n\n{'═'*52}")
    print(f"📋 Busca por planilha | Resumo")
    print(f"{'═'*52}")
    print(f"   Encontrados (match)  : {ok}/{total}")
    print(f"   Fotos coletadas      : {tot_fotos}")
    print(f"   Avaliações coletadas : {tot_rev}")
    print(f"   Banda                : {mb:.1f} MB ({mb/max(metrics['pois'],1):.2f}/POI)")
    print(f"   Tempo                : {dur/60:.1f} min")
    print(f"   💾 {out_json}")
    print(f"{'═'*52}")

    # ── Passo de recuperação IA + Gemini + descoberta (só os reprocessados) ──
    if recuperar:
        _recuperar_ia_e_descoberta(state["results"], target_uf, sessao_nome, rows_reproc, poligono)
        salvar_json(out_json, state["results"])

    return out_json


def _place_ids_no_banco() -> set:
    try:
        import psycopg2
        env = {}
        for ln in (config.BASE_DIR / ".env").read_text(encoding="utf-8").splitlines():
            ln = ln.strip()
            if ln and not ln.startswith("#") and "=" in ln:
                k, v = ln.split("=", 1); env[k] = v.strip().strip('"')
        c = psycopg2.connect(host=env.get("POSTGRES_HOST", "localhost"), port=env.get("POSTGRES_PORT", "5432"),
                             user=env["POSTGRES_USER"], password=env["POSTGRES_PASSWORD"],
                             dbname=env.get("POSTGRES_DB", "comercialradar"))
        cur = c.cursor(); cur.execute("SELECT place_id FROM pois WHERE place_id IS NOT NULL")
        ids = {r[0] for r in cur.fetchall()}; c.close()
        return ids
    except Exception:
        return set()


MAX_VIZ_M = 20000  # candidato vizinho só é confiável a ≤20km da coordenada da planilha


def _cands_confiaveis(r):
    """Candidatos vizinhos DENTRO do portão de distância. JSONs antigos acumulam
    candidatos coletados antes do portão existir (dist_m de milhares de km) —
    eles causavam recuperado_ia/descoberto do outro lado do Brasil."""
    return [c for c in (r.get("candidatos_proximos") or [])
            if c.get("dist_m") is not None and c["dist_m"] <= MAX_VIZ_M]


def _recuperar_ia_e_descoberta(results, target_uf, sessao_nome, rows_alvo, poligono=None):
    import ai_decisor
    # Só opera sobre os itens REPROCESSADOS agora (rows_alvo)
    alvo = [r for r in results if r.get("_row") in rows_alvo]

    # 1) IA (OpenAI) decide equivalência p/ não-matches com candidatos vizinhos
    #    (só candidatos ≤20km — o índice devolvido pela IA referencia a lista FILTRADA)
    pend = [(r, _cands_confiaveis(r)) for r in alvo if not r.get("match_valido")]
    pend = [(r, cands) for r, cands in pend if cands]
    print(f"\n🤖 Recuperação IA (vizinhos): {len(pend)} POIs → consultando OpenAI...")
    if pend:
        consultas = [{"nome": r.get("nome_planilha", ""), "endereco": r.get("endereco_planilha", ""),
                      "candidatos": cands} for r, cands in pend]
        decisoes = ai_decisor.decidir_lote(consultas)
        rec_ia = 0
        for (r, cands), dec in zip(pend, decisoes):
            idx = dec.get("idx")
            if idx is not None and 0 <= idx < len(cands) and dec.get("confianca", 0) >= 0.6:
                c = cands[idx]
                # Cinto e suspensório: nunca aceitar recuperação além do portão
                if c.get("dist_m") is None or c["dist_m"] > MAX_VIZ_M:
                    continue
                r.update({
                    "nome": c.get("nome", ""), "endereco": c.get("endereco", ""),
                    "maps_lat": c.get("lat"), "maps_lng": c.get("lng"),
                    "maps_url": c.get("maps_url", ""), "place_id": c.get("place_id"),
                    "distancia_m": c.get("dist_m"), "status": "recuperado_ia",
                    "match_valido": True, "ia_confianca": dec.get("confianca"),
                    "ia_motivo": dec.get("motivo"),
                })
                if area_utils.gate_registro(r, poligono):
                    rec_ia += 1
        print(f"   ✅ Recuperados pela IA (vizinhos): {rec_ia}")

    # 2) 3ª camada: Gemini + Google Search grounding (batch) p/ os AINDA sem match
    ainda = [r for r in alvo if not r.get("match_valido") and r.get("nome_planilha")]
    if ainda:
        try:
            import gemini_localizador
            print(f"\n🔷 Gemini localização (grounding, batch): {len(ainda)} POIs...")
            consultas = [{"nome": r["nome_planilha"],
                          "municipio": r.get("endereco_planilha", "") or "",
                          "uf": target_uf or ""} for r in ainda]
            locs = gemini_localizador.localizar_lote(consultas)
            rec_g = 0
            for r, loc in zip(ainda, locs):
                if not loc or not loc.get("encontrado"):
                    continue
                # Coord do Gemini só vale se ≤40km da coord da planilha (mesma regra
                # do --gemini-direto); senão fica a coordenada da própria planilha.
                la, lo = r.get("lat_origem"), r.get("lng_origem")
                gla, glo = loc.get("lat"), loc.get("lng")
                if gla is not None and glo is not None:
                    if la is None:
                        la, lo = gla, glo
                    else:
                        try:
                            if haversine(la, lo, gla, glo) <= 40000:
                                la, lo = gla, glo
                        except Exception:
                            pass
                r.update({
                    "nome": loc.get("nome_oficial") or r["nome_planilha"],
                    "endereco": loc.get("endereco", "") or "",
                    "telefone": loc.get("telefone", "") or "",
                    "website": loc.get("website", "") or "",
                    "categoria": loc.get("categoria", "") or "",
                    "avaliacao": loc.get("avaliacao"),
                    "total_avaliacoes": loc.get("total_avaliacoes"),
                    "status_horario": loc.get("horario", "") or "",
                    "preco_medio": loc.get("preco_medio", "") or "",
                    "maps_lat": la, "maps_lng": lo,
                    "status": "recuperado_gemini", "match_valido": True, "fonte_dado": "gemini",
                    "ia_resposta": json.dumps(loc, ensure_ascii=False),
                })
                if area_utils.gate_registro(r, poligono):
                    rec_g += 1
            print(f"   🔷 Recuperados pelo Gemini: {rec_g}")
        except Exception as e:
            print(f"   ⚠️  Gemini falhou: {str(e)[:80]}")

    # 3) Descoberta: candidatos vizinhos novos (place_id não no banco) → base
    ja = _place_ids_no_banco()
    ja |= {r.get("place_id") for r in results if r.get("place_id")}
    novos = {}
    for r in alvo:
        for c in _cands_confiaveis(r):
            pid = c.get("place_id")
            if not pid or pid in ja or pid in novos:
                continue
            # Descoberto fora do polígono da área não interessa (nem entra no JSON)
            if poligono and c.get("lat") is not None and \
               not area_utils.ponto_no_poligono(c.get("lat"), c.get("lng"), poligono):
                continue
            novos[pid] = {
                "fonte": "descoberto", "sessao": sessao_nome,
                "nome": c.get("nome", ""), "endereco": c.get("endereco", ""),
                "maps_lat": c.get("lat"), "maps_lng": c.get("lng"),
                "lat_origem": c.get("lat"), "lng_origem": c.get("lng"),
                "maps_url": c.get("maps_url", ""), "place_id": pid,
                "status": "descoberto", "match_valido": True,
                "fotos": [], "comentarios": [], "horarios": {},
            }
    if novos:
        results.extend(novos.values())
        print(f"   🔎 Locais novos descobertos (enriquecem a base): {len(novos)}")


def _gemini_direto(pendentes, results_existentes, target_uf, sessao_nome, out_json, poligono=None):
    """Resíduo (já falhou no Maps) → direto pro Gemini grounding, sem browser.
    Paralelo (vários lotes ao mesmo tempo) + progresso + salvamento incremental."""
    import time as _t
    import threading
    from concurrent.futures import ThreadPoolExecutor, as_completed
    import gemini_localizador

    TAM = gemini_localizador.TAM_LOTE
    blocos = [pendentes[i:i + TAM] for i in range(0, len(pendentes), TAM)]
    print(f"\n🔷 Modo Gemini-direto: {len(pendentes)} POIs em {len(blocos)} lotes → grounding paralelo (sem Maps)...\n")

    def _monta(it, loc):
        reg = {
            "fonte": "planilha", "sessao": sessao_nome, "_row": it.get("_row"),
            "nome_planilha": it["nome"], "endereco_planilha": it.get("endereco", ""),
            "lat_origem": it.get("lat"), "lng_origem": it.get("lng"),
            "status": "nao_encontrado", "match_valido": False,
            "fotos": [], "comentarios": [], "horarios": {},
        }
        if loc and loc.get("encontrado"):
            la, lo = it.get("lat"), it.get("lng")
            gla, glo = loc.get("lat"), loc.get("lng")
            if gla is not None and glo is not None and it.get("lat") is not None:
                try:
                    if haversine(it["lat"], it["lng"], gla, glo) <= 40000:
                        la, lo = gla, glo
                except Exception:
                    pass
            reg.update({
                "nome": loc.get("nome_oficial") or it["nome"],
                "endereco": loc.get("endereco", "") or "",
                "telefone": loc.get("telefone", "") or "",
                "website": loc.get("website", "") or "",
                "categoria": loc.get("categoria", "") or "",
                "avaliacao": loc.get("avaliacao"), "total_avaliacoes": loc.get("total_avaliacoes"),
                "status_horario": loc.get("horario", "") or "",
                "preco_medio": loc.get("preco_medio", "") or "",
                "maps_lat": la, "maps_lng": lo,
                "status": "recuperado_gemini", "match_valido": True, "fonte_dado": "gemini",
                "ia_resposta": json.dumps(loc, ensure_ascii=False),
            })
        area_utils.gate_registro(reg, poligono)
        return reg

    def _proc(bloco):
        consultas = [{"nome": it["nome"], "municipio": it.get("endereco", "") or "",
                      "uf": (it.get("uf") or target_uf or "")} for it in bloco]
        locs = gemini_localizador.localizar_bloco(consultas)
        return [_monta(it, loc) for it, loc in zip(bloco, locs)]

    novos, rec, done = [], 0, 0
    lock = threading.Lock()
    ini = _t.time()
    with ThreadPoolExecutor(max_workers=8) as ex:
        futs = [ex.submit(_proc, b) for b in blocos]
        for k, fut in enumerate(as_completed(futs), 1):
            try:
                regs = fut.result()
            except Exception:
                regs = []
            with lock:
                novos.extend(regs)
                rec += sum(1 for r in regs if r.get("match_valido"))
                done += len(regs)
                if k % 5 == 0 or k == len(blocos):
                    salvar_json(out_json, results_existentes + novos)
                print(f"\r🔷 lotes {k}/{len(blocos)} | POIs {done}/{len(pendentes)} | "
                      f"recuperados {rec} | {(_t.time()-ini)/60:.1f}min", flush=True)

    salvar_json(out_json, results_existentes + novos)
    print(f"\n\n{'═'*52}")
    print(f"🔷 Gemini-direto | Resumo")
    print(f"{'═'*52}")
    print(f"   Recuperados pelo Gemini : {rec}/{len(pendentes)}")
    print(f"   Ainda sem match         : {len(pendentes) - rec}")
    print(f"   Tempo                   : {(_t.time()-ini)/60:.1f} min")
    print(f"   💾 {out_json}")
    print(f"{'═'*52}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("planilha", help="Caminho para .xlsx ou .csv")
    parser.add_argument("--workers", type=int, default=6)
    parser.add_argument("--cidade", default="", help="Cidade p/ enviesar quando a linha não tem endereço")
    parser.add_argument("--no-proxy", action="store_true", help="Roda direto, sem gastar banda de proxy (teste)")
    parser.add_argument("--ingest", action="store_true", help="Após coletar, ingere no Postgres via Prisma")
    parser.add_argument("--limit", type=int, default=0, help="Processa só as N primeiras linhas (teste)")
    parser.add_argument("--retry-failed", action="store_true",
                        help="Reprocessa só as linhas não-'ok' (divergente/nao_encontrado/erro)")
    parser.add_argument("--recuperar", action="store_true",
                        help="Recuperação por vizinhança ('próximo daqui') + decisão por IA + descoberta")
    parser.add_argument("--headful", action="store_true",
                        help="Abre o navegador VISÍVEL (não-headless) com 1 worker, p/ acompanhar o processo")
    parser.add_argument("--gemini-direto", action="store_true",
                        help="Resíduo que já falhou no Maps → direto pro Gemini (sem browser/Maps). Use com --retry-failed.")
    parser.add_argument("--area", default="",
                        help="JSON com o polígono da área válida ({'polygon': [[lat,lng],...]}). "
                             "POI fora do polígono → status 'fora_da_area' (não ingere).")
    args = parser.parse_args()

    sheet = Path(args.planilha)
    if not sheet.exists():
        print(f"Arquivo não encontrado: {sheet}")
        return

    workers = args.workers
    headless = True
    if args.headful:
        workers = 1          # 1 janela visível só
        headless = False
        print("👁️  Modo VISÍVEL: 1 worker, navegador aberto para acompanhamento.")

    out_json = asyncio.run(run(sheet, workers, args.cidade, not args.no_proxy,
                               limit=args.limit, retry_failed=args.retry_failed,
                               recuperar=args.recuperar, headless=headless,
                               gemini_direto=args.gemini_direto,
                               area_path=(Path(args.area) if args.area else None)))

    if args.ingest and out_json:
        print("\n▶ Ingerindo no PostgreSQL via Prisma...")
        subprocess.run(["npx", "ts-node", "src/ingest.ts", str(out_json)], shell=True)

    # Geração do mapa HTML — passo padrão de todo run
    try:
        import gerar_mapa_html
        gerar_mapa_html.gerar("todos", "mapa_pois.html")
    except Exception as e:
        print(f"⚠️  Mapa não gerado: {e}")


if __name__ == "__main__":
    main()
